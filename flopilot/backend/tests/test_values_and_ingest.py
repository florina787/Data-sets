"""Value parsing (dates/locale), upload validation and CSV handling."""
from __future__ import annotations

import io
import zipfile
from datetime import date

import openpyxl
import pytest

from app.common import AppError
from app.workbook import values as V
from app.workbook.ingest import validate_upload
from app.workbook.model import csv_to_xlsx, load_model
from tests.conftest import upload


@pytest.mark.parametrize("text,day_first,expected,ambiguous", [
    ("2025-03-04", None, date(2025, 3, 4), False),
    ("15/03/2025", None, date(2025, 3, 15), False),     # only day-first is valid
    ("03/15/2025", None, date(2025, 3, 15), False),     # only month-first is valid
    ("04/05/2025", None, date(2025, 4, 5), True),       # ambiguous -> flagged
    ("04/05/2025", True, date(2025, 5, 4), False),
    ("04.05.2025", True, date(2025, 5, 4), False),
    ("15-Mar-2025", None, date(2025, 3, 15), False),
    ("March 15, 2025", None, date(2025, 3, 15), False),
    ("31/02/2025", None, None, False),                  # impossible date
])
def test_date_parsing(text, day_first, expected, ambiguous):
    assert V.parse_date_text(text, day_first) == (expected, ambiguous)


@pytest.mark.parametrize("text,expected", [
    ("1,234.50", 1234.5), ("$1,200", 1200), ("(50)", -50), ("12%", 0.12), ("€3", 3), ("abc", None), ("1.2.3", None),
])
def test_numeric_text(text, expected):
    assert V.parse_numeric_text(text) == expected


def test_display_formats():
    assert V.format_display(1234.5, '"$"#,##0.00') == "$1,234.50"
    assert V.format_display(0.256, "0.0%") == "25.6%"
    assert V.format_display(-5, "#,##0;(#,##0)") == "(5)"
    assert V.format_display(3.0, "General") == "3"


def test_rejects_empty_unsupported_and_malformed():
    with pytest.raises(AppError) as e:
        validate_upload("a.xlsx", b"")
    assert e.value.code == "empty_file"
    for ext in (".xls", ".xlsm", ".xlsb", ".ods"):
        with pytest.raises(AppError) as e:
            validate_upload(f"a{ext}", b"PK\x03\x04data")
        assert e.value.code == "unsupported_format" and e.value.status == 415
    with pytest.raises(AppError) as e:
        validate_upload("a.xlsx", b"this is not a zip")
    assert e.value.code == "malformed_file"
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as z:
        z.writestr("hello.txt", "x")
    with pytest.raises(AppError) as e:
        validate_upload("a.xlsx", buf.getvalue())
    assert e.value.code == "malformed_file"


def test_password_protected_detected():
    ole = b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" + b"\x00" * 600
    with pytest.raises(AppError) as e:
        validate_upload("secret.xlsx", ole)
    assert e.value.code == "password_protected_or_legacy"
    assert "password" in e.value.message.lower()


def test_macro_package_rejected_even_when_named_xlsx():
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as z:
        z.writestr("xl/workbook.xml", "<workbook/>")
        z.writestr("xl/vbaProject.bin", b"\x00")
    with pytest.raises(AppError) as e:
        validate_upload("renamed.xlsx", buf.getvalue())
    assert e.value.code == "macros_present"


def test_size_limit(settings_override):
    settings_override(max_upload_mb=0)
    with pytest.raises(AppError) as e:
        validate_upload("a.csv", b"a,b\n1,2\n")
    assert e.value.code == "file_too_large"


def test_csv_conversion_neutralises_formula_text():
    data = b"Name,Amount,Note\nAlice,10,=HYPERLINK(\"http://evil\",\"x\")\nBob,2.5,+cmd\n"
    m = load_model(csv_to_xlsx(data, "Data"))
    ws = m.ws_f("Data")
    assert ws["C2"].data_type == "s" and ws["C2"].value.startswith("=HYPERLINK")
    assert ws["B2"].value == 10 and ws["B3"].value == 2.5
    assert ws.freeze_panes == "A2"


def test_csv_upload_keeps_original(client, H):
    wb = upload(client, H, b"Region;Amount\nNorth;5\nSouth;7\n", "regional.csv")
    assert wb["source_format"] == "csv"
    versions = client.get(f"/api/workbooks/{wb['id']}/versions", headers=H).json()
    assert versions[0]["kind"] == "converted_from_csv"
    profile = client.get(f"/api/workbooks/{wb['id']}/profile", headers=H).json()
    assert profile["sheets"][0]["table"]["row_count"] == 2


def test_malformed_upload_returns_safe_error(client, H):
    r = client.post("/api/workbooks", headers=H, files={"file": ("bad.xlsx", b"PK\x03\x04garbage")})
    assert r.status_code == 400
    body = r.json()
    assert body["code"] == "malformed_file" and "Traceback" not in r.text


def test_empty_workbook_profiles_cleanly(client, H):
    wb = openpyxl.Workbook()
    created = upload(client, H, wb, "empty.xlsx")
    profile = client.get(f"/api/workbooks/{created['id']}/profile", headers=H).json()
    assert profile["sheets"][0]["table"] is None
    assert any(i["kind"] == "empty_sheet" for i in profile["issues"])
