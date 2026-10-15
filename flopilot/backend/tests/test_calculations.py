"""Deterministic calculations against small synthetic workbooks with hand-computed expectations."""
from __future__ import annotations

import io
from datetime import datetime

import openpyxl
import pytest

from app.common import AppError
from app.workbook import query
from app.workbook.model import load_model


def make(rows: list[list], headers: list[str], sheet: str = "Data", extra=None) -> bytes:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = sheet
    ws.append(headers)
    for r in rows:
        ws.append(r)
    if extra:
        extra(wb)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


@pytest.fixture()
def sales():
    rows = [
        ["North", "A", 100, datetime(2025, 1, 5)],
        ["North", "B", 50, datetime(2025, 1, 20)],
        ["South", "A", 200, datetime(2025, 2, 3)],
        ["South", "B", None, datetime(2025, 2, 10)],        # blank measure
        ["South", "A", "1,000", datetime(2025, 4, 1)],       # number stored as text
        ["West", "B", "n/a", "2025-04-02"],                  # text in measure, date as text
        ["north ", "A", 25, datetime(2025, 4, 9)],           # spelling variant (case/space)
    ]
    return load_model(make(rows, ["Region", "Product", "Amount", "Date"]), "wb", 1)


def test_sum_by_group_with_exclusions(sales):
    r = query.aggregate(sales, "Data", "Amount", "sum", ["Region"])
    by = {row["group"][0]: row["value"] for row in r["rows"]}
    # "North" and "north " are grouped case/space-insensitively
    assert by == {"North": 175.0, "South": 200.0, "West": None}
    assert r["total"] == 375.0
    assert r["excluded"] == {"blank": 1, "text": 1, "numeric_text": 1, "error": 0}
    assert any("stored as text" in w for w in r["warnings"])
    assert r["sources"][0]["range"] == "C2:C8" and r["sources"][0]["header"] == "Amount"


def test_average_and_count_and_filters(sales):
    avg = query.aggregate(sales, "Data", "Amount", "average", [], [{"column": "Product", "op": "eq", "value": "a"}])
    assert avg["rows"][0]["value"] == pytest.approx((100 + 200 + 25) / 3)
    assert avg["filters"] == [{"column": "Product", "op": "eq", "value": "a"}]
    count = query.aggregate(sales, "Data", None, "count", ["Product"])
    assert {r["group"][0]: r["value"] for r in count["rows"]} == {"A": 4, "B": 3}
    gt = query.aggregate(sales, "Data", "Amount", "sum", [], [{"column": "Amount", "op": "gt", "value": 60}])
    assert gt["rows"][0]["value"] == 300.0


def test_missing_column_and_sheet_errors(sales):
    with pytest.raises(AppError) as e:
        query.aggregate(sales, "Data", "Revenue", "sum")
    assert e.value.code == "column_not_found" and "Available columns" in e.value.message
    with pytest.raises(AppError) as e:
        query.aggregate(sales, "Nope", "Amount", "sum")
    assert e.value.code == "sheet_not_found"


def test_time_series_gaps_and_text_dates(sales):
    ts = query.time_series(sales, "Data", "Date", "Amount")
    assert ts["periods"] == ["2025-01", "2025-02", "2025-03", "2025-04"]
    values = {p["period"]: p["value"] for p in ts["series"]["All"]}
    assert values == {"2025-01": 150.0, "2025-02": 200.0, "2025-03": None, "2025-04": 25.0}
    assert ts["undated_rows"] == 1
    assert any("gaps, not zeros" in w for w in ts["warnings"])


def test_period_change_and_variance():
    rows = [["X", datetime(2025, 5, 1), 100, 90], ["X", datetime(2025, 6, 1), 40, 50],
            ["Y", datetime(2025, 5, 2), 10, 10], ["Y", datetime(2025, 6, 2), 30, 20]]
    m = load_model(make(rows, ["Item", "Date", "Actual", "Budget"]), "wb", 1)
    pc = query.period_change(m, "Data", "Date", "Actual", "Item")
    assert (pc["previous_period"], pc["current_period"]) == ("2025-05", "2025-06")
    assert pc["rows"][0]["group"] == ["X"] and pc["rows"][0]["value"] == -60.0 and pc["rows"][0]["pct"] == -60.0
    var = query.compare_measures(m, "Data", "Actual", "Budget", "Item")
    by = {r["group"][0]: r for r in var["rows"]}
    assert by["X"]["value"] == 0.0 and by["Y"]["value"] == 10.0 and by["Y"]["pct"] == 33.33
    assert var["total"] == {"actual": 180.0, "budget": 170.0}


def test_correlation_is_labelled_as_association():
    rows = [[i, 2 * i + 1] for i in range(1, 11)]
    m = load_model(make(rows, ["X", "Y"]), "wb", 1)
    c = query.correlation(m, "Data", "X", "Y")
    assert c["value"] == pytest.approx(1.0)
    assert any("causation" in w for w in c["warnings"])


def test_formula_cells_use_cached_values_not_expressions(client, H):
    from tests.conftest import DEMO

    m = load_model(DEMO.read_bytes(), "wb", 1)
    table = m.table("Sales")
    profit = table.column("Profit")
    # Profit is a formula column (=K-L); values come from the saved results
    assert (2, profit.index) in table.formulas
    r = query.aggregate(m, "Sales", "Profit", "sum")
    net = query.aggregate(m, "Sales", "Net Revenue", "sum")["rows"][0]["value"]
    cost = query.aggregate(m, "Sales", "Cost", "sum")["rows"][0]["value"]
    typed_row_40 = 1234.56
    k40, l40 = m.ws_v("Sales")["K40"].value, m.ws_v("Sales")["L40"].value
    assert r["rows"][0]["value"] == pytest.approx(net - cost - (k40 - l40) + typed_row_40)


def test_empty_sheet_is_reported():
    def extra(wb):
        wb.create_sheet("Empty")
    m = load_model(make([[1]], ["A"], extra=extra), "wb", 1)
    with pytest.raises(AppError) as e:
        query.aggregate(m, "Empty", "A", "sum")
    assert e.value.code == "empty_sheet"


def test_ambiguous_column_detected():
    m = load_model(make([[1, 2]], ["Gross Revenue", "Net Revenue"]), "wb", 1)
    with pytest.raises(AppError) as e:
        query.aggregate(m, "Data", "revenue", "sum")
    assert e.value.code == "ambiguous_column"
    assert set(e.value.details["candidates"]) == {"Gross Revenue", "Net Revenue"}


def test_mixed_currency_units_are_flagged():
    def extra(wb):
        ws = wb["Data"]
        ws["B2"].number_format = '"$"#,##0'
        ws["B3"].number_format = '"€"#,##0'
    m = load_model(make([["a", 10], ["b", 20]], ["Item", "Amount"], extra=extra), "wb", 1)
    r = query.aggregate(m, "Data", "Amount", "sum")
    assert r["rows"][0]["value"] == 30.0
    assert any("mixes currency formats" in w for w in r["warnings"])
