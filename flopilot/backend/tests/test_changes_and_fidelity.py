"""Change application, validation gates, formatting preservation and fidelity detection."""
from __future__ import annotations

import io

import openpyxl
import pytest
from openpyxl.chart import BarChart, Reference

from app.common import AppError
from app.workbook import cleaning, formulas
from app.workbook.changes import apply_change_set, build_formula_column, build_summary_sheet, make_change_set, value_change
from app.workbook.fidelity import compare_packages, detect_features
from app.workbook.model import load_model
from app.workbook.query import aggregate
from app.workbook.validate import validate_candidate
from tests.conftest import DEMO


@pytest.fixture(scope="module")
def base():
    return DEMO.read_bytes()


@pytest.fixture(scope="module")
def demo(base):
    return load_model(base, "wb", 1)


def test_formula_column_preserves_formatting(base, demo):
    p = formulas.build_formula_proposal(demo, "Inventory", {"kind": "ratio", "numerator": "On Hand", "denominator": "Reorder Level",
                                                            "as_percent": False, "new_column": "Cover"})
    cs = build_formula_column(demo, p, {})
    out = apply_change_set(base, cs)
    v = validate_candidate(base, out, cs)
    assert v["passed"], [c for c in v["checks"] if not c["passed"]]
    wb = openpyxl.load_workbook(io.BytesIO(out))
    inv = wb["Inventory"]
    assert inv["H1"].value == "Cover" and inv["H2"].value == '=IFERROR(D2/E2,"")'
    assert inv.tables["InventoryTable"].ref == "A1:H16"           # Excel table extended to the new column
    assert inv.freeze_panes == "A2"
    assert inv["F2"].number_format == '"$"#,##0.00'               # untouched formats kept
    assert len(inv.data_validations.dataValidation) == 1 and len(list(inv.conditional_formatting)) == 1
    assert wb["Sales"].column_dimensions["D"].width == 13
    assert wb["Sales"].auto_filter.ref == "A1:M121"
    assert wb.calculation.fullCalcOnLoad is True
    # cached results were refreshed, so the new version reads correctly without Excel
    wbv = openpyxl.load_workbook(io.BytesIO(out), data_only=True)
    assert wbv["Inventory"]["H2"].value == pytest.approx(wbv["Inventory"]["D2"].value / wbv["Inventory"]["E2"].value)
    assert wbv["Sales"]["M2"].value is not None and wbv["Summary"]["B4"].value is not None
    assert wbv["Expenses"]["H23"].value == "#DIV/0!"


def test_validation_detects_unplanned_change(base, demo):
    cs = cleaning.normalize_categories(demo, "Sales", "Region")
    out = apply_change_set(base, cs)
    wb = openpyxl.load_workbook(io.BytesIO(out))
    wb["Sales"]["G2"] = 999  # tamper: a change that was not approved
    buf = io.BytesIO()
    wb.save(buf)
    v = validate_candidate(base, buf.getvalue(), cs)
    failed = {c["name"] for c in v["checks"] if not c["passed"]}
    assert {"only_planned_cells_changed", "totals_reconcile"} <= failed
    assert not v["passed"]


def test_validation_detects_formula_overwrite(base, demo):
    cs = make_change_set("test", "t", "t", "Sales", [value_change(demo, "Sales", "B2", "x")])
    out = apply_change_set(base, cs)
    wb = openpyxl.load_workbook(io.BytesIO(out))
    wb["Sales"]["M3"] = 5  # overwrite a formula that was not part of the plan
    buf = io.BytesIO()
    wb.save(buf)
    v = validate_candidate(base, buf.getvalue(), cs)
    assert any(c["name"] == "no_unexpected_formula_overwrites" and not c["passed"] for c in v["checks"])


def test_precondition_rejects_changed_cell(base, demo):
    cs = cleaning.normalize_categories(demo, "Sales", "Region")
    wb = openpyxl.load_workbook(io.BytesIO(base))
    first = cs["changes"][0]["cell"]
    wb["Sales"][first] = "Someone else edited this"
    buf = io.BytesIO()
    wb.save(buf)
    with pytest.raises(AppError) as e:
        apply_change_set(buf.getvalue(), cs)
    assert e.value.code == "precondition_failed"


def test_cleaning_values_are_correct(base, demo):
    cs = cleaning.standardize_dates(demo, "Sales", "Order Date")
    assert cs["details"]["day_first"] is False  # 03/15/2025-style values prove month-first
    out = load_model(apply_change_set(base, cs), "wb", 2)
    table = out.table("Sales")
    assert all(hasattr(v, "year") for v in table.values(table.column("Order Date")))
    assert out.ws_f("Sales")["B32"].number_format == "yyyy-mm-dd"
    regions = cleaning.normalize_categories(demo, "Sales", "Region")
    assert regions["details"]["mapping"] == {"north america": "North America", "Europe ": "Europe", "EUROPE": "Europe",
                                             "Asia Pacfic": "Asia Pacific"}


def test_numeric_text_conversion_changes_totals_exactly(base, demo):
    before = aggregate(demo, "Sales", "Units", "sum")["rows"][0]["value"]
    cs = cleaning.convert_numeric_text(demo, "Sales", "Units")
    out = apply_change_set(base, cs)
    assert validate_candidate(base, out, cs)["passed"]
    after = aggregate(load_model(out, "wb", 2), "Sales", "Units", "sum")["rows"][0]["value"]
    assert after == before + float(demo.ws_f("Sales")["G27"].value)


def test_structural_delete_blocked_when_referenced(demo):
    with pytest.raises(AppError) as e:
        cleaning.duplicates_change(demo, "Sales", remove=True)
    assert e.value.code == "structural_edit_blocked"
    flag = cleaning.duplicates_change(demo, "Sales", remove=False)
    assert flag["action_type"] == "add_flag_column" and len(flag["changes"]) == 1 + 3


def test_delete_rows_on_unreferenced_sheet(base, demo):
    cs = cleaning.duplicates_change(demo, "Customers", remove=True)
    assert cs["structural"]["delete_rows"] == [42]
    out = apply_change_set(base, cs)
    v = validate_candidate(base, out, cs)
    assert v["passed"], [c for c in v["checks"] if not c["passed"]]
    assert openpyxl.load_workbook(io.BytesIO(out))["Customers"].max_row == 41


def test_summary_sheet_with_native_chart(base, demo):
    result = aggregate(demo, "Sales", "Net Revenue", "sum", ["Region"])
    cs = build_summary_sheet(demo, result, "Revenue by region", {"type": "bar", "title": "Revenue by region"})
    out = apply_change_set(base, cs)
    assert validate_candidate(base, out, cs)["passed"]
    wb = openpyxl.load_workbook(io.BytesIO(out))
    assert wb.sheetnames[-1] == "Revenue by region"
    assert detect_features(out)["inventory"].get("charts") == 1


def chart_workbook() -> bytes:
    wb = openpyxl.Workbook()
    ws = wb.active
    for i in range(1, 6):
        ws.append([f"r{i}", i])
    chart = BarChart()
    chart.add_data(Reference(ws, min_col=2, min_row=1, max_row=5))
    ws.add_chart(chart, "D2")
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_fidelity_is_measured_not_assumed():
    data = chart_workbook()
    features = detect_features(data)
    assert features["requires_acknowledgement"] and any(f["feature"] == "charts" for f in features["at_risk"])
    model = load_model(data, "wb", 1)
    cs = make_change_set("test", "t", "t", "Sheet", [value_change(model, "Sheet", "A1", "first")])
    out = apply_change_set(data, cs)
    measured = compare_packages(data, out)
    v_without_ack = validate_candidate(data, out, cs, acknowledged_risks=False)
    if measured["lost"]:
        assert not v_without_ack["passed"]
        assert validate_candidate(data, out, cs, acknowledged_risks=True)["passed"]
    else:
        assert v_without_ack["passed"]
