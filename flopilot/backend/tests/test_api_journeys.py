"""End-to-end journeys through the HTTP API (upload -> ask -> propose -> approve -> download -> reopen)."""
from __future__ import annotations

import io

import openpyxl

from tests.conftest import DEMO, approve, ask, login, upload


def test_full_journey_upload_to_download(client, H):
    wb = upload(client, H, DEMO.read_bytes(), "Quarterly pack.xlsx")
    wid = wb["id"]
    run, conv = ask(client, H, wid, "Summarize revenue by region")
    assert run["state"] == "COMPLETED" and run["result"]["type"] == "clarification"   # Gross vs Net revenue
    run, conv = ask(client, H, wid, "Net Revenue", conv)
    result = run["result"]
    assert result["type"] == "answer" and result["intent"] == "aggregate"
    assert "Net Revenue" in result["text"] and result["sources"][0]["sheet"] == "Sales"
    assert result["chart"]["type"] == "bar" and result["chart"]["web_only"] is True
    assert any("Net Revenue" in a for a in result["assumptions"])            # confirmed rule is shown

    run, conv = ask(client, H, wid, "Add a profit-margin column", conv)
    assert run["state"] == "AWAITING_APPROVAL"
    op = run["result"]["operations"][0]
    assert op["status"] == "PROPOSED" and op["flags"]["inserts_formulas"] == 120 and op["dry_run"]["passed"]
    assert op["changes_sample"][1]["after"]["formula"] == '=IFERROR(M2/K2,"")'
    assert client.get(f"/api/workbooks/{wid}", headers=H).json()["current_version"] == 1  # nothing applied yet

    applied = approve(client, H, op).json()
    assert applied["status"] == "APPLIED" and applied["result_version"] == 2
    assert all(c["passed"] for c in applied["validation"]["checks"])
    assert client.get(f"/api/runs/{run['id']}", headers=H).json()["state"] == "COMPLETED"

    d = client.get(f"/api/workbooks/{wid}/versions/2/download", headers=H)
    assert d.status_code == 200 and "attachment" in d.headers["content-disposition"]
    reopened = openpyxl.load_workbook(io.BytesIO(d.content))
    assert reopened["Sales"]["N1"].value == "Profit Margin"
    assert reopened.sheetnames == ["Sales", "Expenses", "Inventory", "Customers", "Summary", "Notes"]
    original = client.get(f"/api/workbooks/{wid}/versions/1/download", headers=H)
    assert original.content == DEMO.read_bytes()  # original upload is immutable


def test_stale_approval_and_idempotency(client, H, demo_wb):
    run, conv = ask(client, H, demo_wb, "Clean the Sales sheet")
    ops = run["result"]["operations"]
    assert len(ops) >= 3 and all(o["status"] == "PROPOSED" for o in ops)
    first = approve(client, H, ops[0], key="once-1").json()
    assert first["status"] == "APPLIED"
    replay = approve(client, H, ops[0], key="once-1").json()
    assert replay.get("idempotent_replay") and replay["result_version"] == first["result_version"]
    assert len(client.get(f"/api/workbooks/{demo_wb}/versions", headers=H).json()) == 2   # no duplicate write

    stale = approve(client, H, ops[1])
    assert stale.status_code == 409 and stale.json()["code"] == "stale_approval"
    assert client.get(f"/api/operations/{ops[1]['id']}", headers=H).json()["status"] == "STALE"

    pending = client.get(f"/api/operations/{ops[2]['id']}", headers=H).json()
    assert pending["out_of_date"] is True
    refreshed = client.post(f"/api/operations/{ops[2]['id']}/refresh", headers=H).json()
    new_op = refreshed["operations"][0]
    assert new_op["base_version"] == 2
    assert approve(client, H, new_op).json()["status"] == "APPLIED"


def test_hash_mismatch_rejected(client, H, demo_wb):
    run, _ = ask(client, H, demo_wb, "Normalize the Region categories on Sales")
    op = run["result"]["operations"][0]
    r = client.post(f"/api/operations/{op['id']}/approve", headers=H,
                    json={"expected_version": 1, "change_set_hash": "0" * 64, "acknowledge_risks": False})
    assert r.status_code == 409 and r.json()["code"] == "approval_mismatch"


def test_reject_and_cancel_leave_workbook_unchanged(client, H, demo_wb):
    run, _ = ask(client, H, demo_wb, "Standardize the dates in Order Date")
    op = run["result"]["operations"][0]
    assert client.post(f"/api/operations/{op['id']}/reject", headers=H).json()["status"] == "REJECTED"
    assert approve(client, H, op).status_code == 409
    assert client.get(f"/api/runs/{run['id']}", headers=H).json()["state"] == "CANCELLED"
    run, _ = ask(client, H, demo_wb, "Trim extra spaces in Region on Sales")
    assert client.post(f"/api/runs/{run['id']}/cancel", headers=H).json()["state"] == "CANCELLED"
    assert client.get(f"/api/operations/{run['result']['operations'][0]['id']}", headers=H).json()["status"] == "REJECTED"
    assert client.get(f"/api/workbooks/{demo_wb}", headers=H).json()["current_version"] == 1


def test_restore_previous_version(client, H, demo_wb):
    run, conv = ask(client, H, demo_wb, "Fix the formula error in Expenses!H23")
    op = next(o for o in run["result"]["operations"])
    assert op["flags"]["replaces_formulas"] == 1 and op["destructive"]
    assert approve(client, H, op).json()["result_version"] == 2
    run, conv = ask(client, H, demo_wb, "Undo the last change", conv)
    restore_op = run["result"]["operations"][0]
    assert restore_op["action_type"] == "restore_version" and restore_op["target_version"] == 1
    done = approve(client, H, restore_op).json()
    assert done["result_version"] == 3
    v3 = client.get(f"/api/workbooks/{demo_wb}/versions/3/download", headers=H).content
    assert v3 == client.get(f"/api/workbooks/{demo_wb}/versions/1/download", headers=H).content
    kinds = [v["kind"] for v in client.get(f"/api/workbooks/{demo_wb}/versions", headers=H).json()]
    assert kinds == ["restore", "edit", "original"]


def test_report_sheet_journey(client, H, demo_wb):
    run, conv = ask(client, H, demo_wb, "Compare actual spending with budget")
    assert run["result"]["chart"]["series"][0]["name"] == "Amount"
    assert any("Amount is used as actual" in a for a in run["result"]["assumptions"])
    run, conv = ask(client, H, demo_wb, "Add this summary to the workbook as a report sheet", conv)
    op = run["result"]["operations"][0]
    assert op["flags"]["adds_sheet"] and op["new_sheet"]["chart"]["type"] == "bar"
    applied = approve(client, H, op).json()
    wb = openpyxl.load_workbook(io.BytesIO(client.get(f"/api/workbooks/{demo_wb}/versions/{applied['result_version']}/download", headers=H).content))
    assert "Amount report" in wb.sheetnames and len(wb["Amount report"]._charts) == 1


def test_access_control(client, H, demo_wb):
    assert client.get(f"/api/workbooks/{demo_wb}").status_code == 401
    assert client.get(f"/api/workbooks/{demo_wb}", headers={"Authorization": "Bearer nope"}).status_code == 401
    other = login(client, "other")
    for path in (f"/api/workbooks/{demo_wb}", f"/api/workbooks/{demo_wb}/profile", f"/api/workbooks/{demo_wb}/versions/1/download",
                 f"/api/workbooks/{demo_wb}/sheets/Sales/cells"):
        r = client.get(path, headers=other)
        assert r.status_code == 404 and r.json()["code"] == "workbook_not_found"
    r = client.post(f"/api/workbooks/{demo_wb}/chat", headers=other, json={"message": "hi"})
    assert r.status_code == 404
    run, _ = ask(client, H, demo_wb, "Normalize the Region categories on Sales")
    op = run["result"]["operations"][0]
    assert approve(client, other, op).status_code == 404
    viewer = login(client, "viewer")
    r = approve(client, viewer, op)
    assert r.status_code == 403 and r.json()["code"] == "forbidden"
    assert client.get(f"/api/workbooks/{demo_wb}/versions/1/download", headers=viewer).status_code == 403
    assert client.get("/api/admin/audit", headers=H).status_code == 403
    assert client.get("/api/admin/audit", headers=login(client, "admin")).status_code == 200


def test_prompt_injection_in_workbook_is_inert(client, H, demo_wb):
    run, _ = ask(client, H, demo_wb, "What's in this workbook?")
    text = run["result"]["text"]
    assert "Notes** (hidden)" in text
    profile = client.get(f"/api/workbooks/{demo_wb}/profile", headers=H).json()
    flagged = [i for i in profile["issues"] if i["kind"] == "instruction_like_text"]
    assert flagged and flagged[0]["sheet"] == "Notes" and "will not follow" in flagged[0]["message"]
    assert "audit-team@example.com" not in str(run["result"]) and "api key" not in text.lower()
    events = client.get(f"/api/workbooks/{demo_wb}/activity", headers=H).json()["events"]
    assert not any(e["event_type"] == "export" for e in events)


def test_csv_export_neutralises_formula_injection(client, H):
    wb = upload(client, H, b"Name,Note\nA,=1+1\nB,@SUM(A1)\nC,-5\n", "inj.csv")
    csv_text = client.get(f"/api/workbooks/{wb['id']}/versions/1/download?format=csv", headers=H).content.decode("utf-8-sig")
    assert "'=1+1" in csv_text and "'@SUM(A1)" in csv_text
    assert "\nC,-5" in csv_text  # a real negative number is data, not a formula, and is exported unchanged


def test_grid_window_and_profile_endpoints(client, H, demo_wb):
    g = client.get(f"/api/workbooks/{demo_wb}/sheets/Sales/cells?r0=1&r1=3&c0=1&c1=13", headers=H).json()
    assert g["freeze_panes"] == "A2" and g["max_row"] == 121
    header = {c["c"]: c for c in g["rows"][0]["cells"]}
    assert header[1]["v"] == "Order ID" and header[1]["s"]["b"] == 1
    m2 = next(c for c in g["rows"][1]["cells"] if c["c"] == 13)
    assert m2["f"] == "=K2-L2" and m2["v"].startswith("$")
    assert client.get(f"/api/workbooks/{demo_wb}/sheets/Nope/cells", headers=H).status_code == 404


def test_schemas_and_config(client):
    s = client.get("/api/schemas").json()
    assert s["schema_version"] == "1.0" and {"Plan", "Operation", "Approval", "Error", "Result"} <= set(s)
    cfg = client.get("/api/config").json()
    assert cfg["llm"]["mode"] == "deterministic-demo"
    assert "not implemented" in cfg["integrations"]["excel_add_in"]


def test_analysis_summary_export(client, H, demo_wb):
    _, conv = ask(client, H, demo_wb, "Show missing values by column on Customers")
    md = client.get(f"/api/workbooks/{demo_wb}/analysis-summary?conversation_id={conv}", headers=H).text
    assert md.startswith("# FloPilot analysis summary") and "Segment" in md and "Source:" in md


def test_cross_sheet_lookup_via_chat(client, H, demo_wb):
    run, _ = ask(client, H, demo_wb, "Add the customer Segment from Customers to Sales")
    op = run["result"]["operations"][0]
    fp = op["details"]["formula_proposal"]
    assert fp["sheet"] == "Sales" and fp["new_column"] == "Segment"
    assert fp["first_formula"] == '=XLOOKUP(F2,Customers!$A$2:$A$42,Customers!$C$2:$C$42,"Not found")'
    assert any("Microsoft 365" in c for c in fp["compatibility"])
    assert fp["preview_status"] == "not_evaluated"  # XLOOKUP is outside the preview evaluator; nothing is claimed
    assert approve(client, H, op).json()["status"] == "APPLIED"


def test_chat_rate_limit(client, H, demo_wb, settings_override):
    from app.api import routes
    routes._rate.clear()
    settings_override(chat_rate_per_min=2)
    for _ in range(2):
        assert client.post(f"/api/workbooks/{demo_wb}/chat", headers=H, json={"message": "Show missing values by column"}).status_code == 200
    r = client.post(f"/api/workbooks/{demo_wb}/chat", headers=H, json={"message": "Show missing values by column"})
    assert r.status_code == 429 and r.json()["code"] == "rate_limited" and r.json()["retryable"] is True
    routes._rate.clear()


def test_delete_purges_content_but_keeps_audit(client, H):
    wb = upload(client, H, DEMO.read_bytes(), "to-delete.xlsx")
    run, conv = ask(client, H, wb["id"], "Normalize the Region categories on Sales")
    op_id = run["result"]["operations"][0]["id"]
    assert client.delete(f"/api/workbooks/{wb['id']}", headers=H).json() == {"deleted": True}
    assert client.get(f"/api/workbooks/{wb['id']}", headers=H).status_code == 404
    from app import db
    assert db.fetchone("SELECT COUNT(*) AS n FROM messages WHERE conversation_id = ?", (conv,))["n"] == 0
    assert db.loads(db.fetchone("SELECT payload FROM operations WHERE id = ?", (op_id,))["payload"]) == {"purged": True}
    assert db.fetchone("SELECT COUNT(*) AS n FROM audit_events WHERE workbook_id = ? AND event_type = 'workbook_deleted'", (wb["id"],))["n"] == 1
