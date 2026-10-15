"""Formula intelligence: explanation, error diagnosis, consistency checks, safe proposals.

Nothing here claims a formula *was calculated*. Proposals carry a preview computed by the
limited evaluator (labelled as such) and, when the LibreOffice engine is available, a separate
verification step recalculates the candidate workbook for real (see recalc.py).
"""
from __future__ import annotations

import re
from collections import Counter

from openpyxl.formula.translate import Translator
from openpyxl.utils import get_column_letter
from openpyxl.utils.cell import coordinate_from_string, column_index_from_string

from ..common import AppError
from . import values as V
from .evaluator import _Parser, Unsupported, evaluate_formula, parse_ref, referenced_functions, referenced_ranges
from .model import SheetTable, WorkbookModel
from .query import require_column, require_table

FUNCTION_DOCS: dict[str, tuple[str, str]] = {
    # name: (plain-English description, minimum Excel version)
    "SUM": ("adds numbers", "all"), "AVERAGE": ("averages numbers", "all"), "MIN": ("smallest value", "all"),
    "MAX": ("largest value", "all"), "COUNT": ("counts numeric cells", "all"), "COUNTA": ("counts non-empty cells", "all"),
    "IF": ("chooses between two results based on a condition", "all"),
    "IFERROR": ("returns a fallback when the calculation produces an error", "all"),
    "IFNA": ("returns a fallback when a lookup finds nothing (#N/A)", "2013"),
    "ROUND": ("rounds to a number of decimal places", "all"), "ABS": ("absolute value", "all"),
    "SUMIF": ("adds values that meet one condition", "all"), "SUMIFS": ("adds values that meet several conditions", "2007"),
    "COUNTIF": ("counts cells that meet a condition", "all"), "COUNTIFS": ("counts rows meeting several conditions", "2007"),
    "AVERAGEIFS": ("averages values meeting several conditions", "2007"),
    "VLOOKUP": ("looks up a value in the first column of a table and returns a value from another column", "all"),
    "HLOOKUP": ("looks up a value in the first row of a table", "all"),
    "INDEX": ("returns the value at a position in a range", "all"), "MATCH": ("finds the position of a value in a range", "all"),
    "XLOOKUP": ("looks up a value and returns the matching item from another range", "365/2021"),
    "XMATCH": ("finds the position of a value (modern MATCH)", "365/2021"),
    "FILTER": ("returns rows that meet conditions (spills)", "365/2021"), "UNIQUE": ("returns distinct values (spills)", "365/2021"),
    "SORT": ("sorts a range (spills)", "365/2021"), "SORTBY": ("sorts by another range (spills)", "365/2021"),
    "SEQUENCE": ("generates a list of numbers (spills)", "365/2021"), "LET": ("names intermediate results", "365/2021"),
    "LAMBDA": ("defines a reusable custom function", "365"), "TEXTSPLIT": ("splits text into cells", "365"),
    "VSTACK": ("stacks arrays vertically", "365"), "HSTACK": ("stacks arrays horizontally", "365"),
    "TEXTJOIN": ("joins text with a delimiter", "2019"), "CONCAT": ("joins text", "2019"), "IFS": ("checks several conditions", "2019"),
    "SWITCH": ("matches a value against cases", "2019"), "MAXIFS": ("largest value meeting conditions", "2019"),
    "MINIFS": ("smallest value meeting conditions", "2019"), "CONCATENATE": ("joins text", "all"),
    "YEAR": ("year of a date", "all"), "MONTH": ("month of a date", "all"), "DAY": ("day of a date", "all"),
    "TEXT": ("formats a value as text", "all"), "DATE": ("builds a date", "all"), "EOMONTH": ("end of month", "all"),
    "TODAY": ("today's date (volatile)", "all"), "NOW": ("current date and time (volatile)", "all"),
    "LEN": ("length of text", "all"), "TRIM": ("removes extra spaces", "all"), "UPPER": ("upper-case text", "all"),
    "LOWER": ("lower-case text", "all"), "LEFT": ("leftmost characters", "all"), "RIGHT": ("rightmost characters", "all"),
    "MID": ("characters from the middle", "all"), "AND": ("true if all conditions are true", "all"),
    "OR": ("true if any condition is true", "all"), "NOT": ("reverses a condition", "all"),
    "ISBLANK": ("checks for an empty cell", "all"), "OFFSET": ("reference shifted from a start cell (volatile)", "all"),
    "INDIRECT": ("reference built from text (volatile)", "all"),
}
MODERN = {name for name, (_, ver) in FUNCTION_DOCS.items() if ver in ("365/2021", "365")}
VOLATILE = {"TODAY", "NOW", "RAND", "RANDBETWEEN", "OFFSET", "INDIRECT", "CELL", "INFO"}
FORBIDDEN = {"WEBSERVICE", "HYPERLINK", "CALL", "REGISTER", "REGISTER.ID", "RTD", "EXEC", "FILTERXML", "IMAGE"}

ERROR_HELP = {
    "#DIV/0!": "The formula divides by zero or by an empty cell.",
    "#N/A": "A lookup did not find a matching value.",
    "#REF!": "The formula refers to a cell that no longer exists (e.g. a deleted row, column or sheet).",
    "#NAME?": "Excel does not recognise a function or name in the formula (misspelling, or a function your Excel version lacks).",
    "#VALUE!": "The formula uses the wrong type of value, for example text in arithmetic.",
    "#NUM!": "A numeric calculation is invalid (e.g. square root of a negative number).",
    "#NULL!": "Two ranges were intersected but do not overlap.",
    "#SPILL!": "A dynamic-array result cannot spill because cells in its way are not empty.",
}


# --------------------------------------------------------------------------- policy
def validate_formula(formula: str, model: WorkbookModel | None = None) -> list[str]:
    """Raise if the formula is unsafe or malformed. Returns non-blocking notes."""
    if not formula.startswith("="):
        raise AppError("invalid_formula", "Formulas must start with '='.")
    if len(formula) > 2000:
        raise AppError("invalid_formula", "Formula is too long.")
    try:
        _Parser(formula).parse()
    except Unsupported as exc:
        raise AppError("invalid_formula", f"The formula could not be parsed: {exc}.") from exc
    except Exception as exc:  # tokenizer errors
        raise AppError("invalid_formula", "The formula could not be parsed.") from exc
    notes = []
    funcs = referenced_functions(formula)
    blocked = sorted(set(funcs) & FORBIDDEN)
    if blocked:
        raise AppError("formula_not_allowed", f"Functions that reach external services or run code are not allowed: {', '.join(blocked)}.", 422)
    for ref in referenced_ranges(formula):
        if "[" in ref:
            raise AppError("formula_not_allowed", "Formulas that reference other workbooks are not allowed.", 422)
        if "!" in ref and model is not None:
            sheet = ref.rsplit("!", 1)[0].strip("'")
            if model.resolve_sheet(sheet) is None:
                raise AppError("invalid_formula", f"The formula refers to a sheet that does not exist: '{sheet}'.", 422)
    if set(funcs) & VOLATILE:
        notes.append(f"Uses volatile function(s) {', '.join(sorted(set(funcs) & VOLATILE))}; results change on every recalculation.")
    return notes


def compatibility(formula: str) -> list[str]:
    funcs = set(referenced_functions(formula))
    notes = []
    modern = sorted(funcs & MODERN)
    if modern:
        notes.append(f"{', '.join(modern)} require Excel for Microsoft 365 or Excel 2021+. Older versions show #NAME?.")
    v2019 = sorted(f for f in funcs if FUNCTION_DOCS.get(f, ("", ""))[1] == "2019")
    if v2019:
        notes.append(f"{', '.join(v2019)} require Excel 2019 or later.")
    unknown = sorted(f for f in funcs if f not in FUNCTION_DOCS)
    if unknown:
        notes.append(f"FloPilot has no description for: {', '.join(unknown)}.")
    if not notes:
        notes.append("Uses functions available in all current Excel versions.")
    return notes


# --------------------------------------------------------------------------- explanation
def _header_for(model: WorkbookModel | None, sheet: str, ref: str) -> str | None:
    if model is None:
        return None
    try:
        s, r1, c1, r2, c2 = parse_ref(ref, sheet)
    except Exception:
        return None
    resolved = model.resolve_sheet(s)
    if resolved is None:
        return None
    table = model.table(resolved)
    if table is None or c1 != c2:
        return None
    for col in table.columns:
        if col.index == c1:
            return col.name
    return None


def _describe(node, model, sheet) -> str:
    kind = node[0]
    if kind == "num":
        v = node[1]
        return str(int(v)) if float(v).is_integer() else str(v)
    if kind == "str":
        return f'"{node[1]}"' if node[1] else "an empty result"
    if kind == "bool":
        return "TRUE" if node[1] else "FALSE"
    if kind == "err":
        return str(node[1])
    if kind == "ref":
        header = _header_for(model, sheet, node[1])
        where = node[1]
        if ":" in node[1]:
            return f"the range {where}" + (f" ({header})" if header else "")
        return f"{header} ({where})" if header else where
    if kind == "neg":
        return f"minus {_describe(node[1], model, sheet)}"
    if kind == "bin":
        op, a, b = node[1], _describe(node[2], model, sheet), _describe(node[3], model, sheet)
        words = {"+": "plus", "-": "minus", "*": "times", "/": "divided by", "^": "to the power of", "&": "joined with",
                 "=": "equals", "<>": "is not equal to", "<": "is less than", ">": "is greater than",
                 "<=": "is at most", ">=": "is at least"}
        return f"{a} {words[op]} {b}"
    if kind == "fn":
        name, args = node[1], [_describe(a, model, sheet) for a in node[2]]
        if name == "IFERROR" and len(args) == 2:
            return f"{args[0]}, or {args[1]} if that produces an error"
        if name == "IF" and len(args) >= 2:
            return f"if {args[0]}, then {args[1]}" + (f", otherwise {args[2]}" if len(args) > 2 else "")
        if name in ("SUM", "AVERAGE", "MIN", "MAX", "COUNT", "COUNTA"):
            label = {"SUM": "the sum of", "AVERAGE": "the average of", "MIN": "the smallest of", "MAX": "the largest of",
                     "COUNT": "the count of numbers in", "COUNTA": "the count of non-empty cells in"}[name]
            return f"{label} {', '.join(args)}"
        if name == "ROUND" and len(args) == 2:
            return f"{args[0]} rounded to {args[1]} decimal place(s)"
        if name in ("XLOOKUP",) and len(args) >= 3:
            return f"the value from {args[2]} where {args[1]} matches {args[0]}" + (f" (or {args[3]} if not found)" if len(args) > 3 else "")
        if name == "VLOOKUP" and len(args) >= 3:
            return f"column {args[2]} of {args[1]} for the row whose first column matches {args[0]}"
        if name == "INDEX" and len(args) >= 2:
            return f"the item at position {args[1]} in {args[0]}"
        if name == "MATCH" and len(args) >= 2:
            return f"the position of {args[0]} within {args[1]}"
        if name in ("SUMIFS",) and len(args) >= 3:
            conds = [f"{args[i]} = {args[i + 1]}" for i in range(1, len(args) - 1, 2)]
            return f"the sum of {args[0]} where {' and '.join(conds)}"
        return f"{name}({', '.join(args)})"
    return "?"


def explain_formula(formula: str, model: WorkbookModel | None = None, sheet: str = "", cell: str | None = None) -> dict:
    if not formula.startswith("="):
        formula = "=" + formula
    try:
        ast = _Parser(formula).parse()
        plain = _describe(ast, model, sheet)
        parse_ok = True
    except Exception:
        plain, parse_ok = "FloPilot could not parse this formula; it may use syntax outside the supported subset.", False
    refs = []
    for ref in referenced_ranges(formula):
        absolute = "$" in ref
        header = _header_for(model, sheet, ref)
        refs.append({
            "ref": ref,
            "header": header,
            "cross_sheet": "!" in ref,
            "external": "[" in ref,
            "anchoring": ("absolute (stays fixed when copied)" if ref.replace("!", "").count("$") >= 2
                          else "mixed (partly fixed)" if absolute else "relative (shifts when copied)"),
        })
    funcs = [{"name": f, "description": FUNCTION_DOCS.get(f, ("not in FloPilot's reference", ""))[0],
              "min_version": FUNCTION_DOCS.get(f, ("", "unknown"))[1]} for f in referenced_functions(formula)]
    return {
        "kind": "formula_explanation", "formula": formula, "cell": cell, "sheet": sheet,
        "plain_english": (f"{cell} = " if cell else "") + plain if parse_ok else plain,
        "functions": funcs, "references": refs, "compatibility": compatibility(formula),
        "volatile": sorted(set(referenced_functions(formula)) & VOLATILE),
    }


# --------------------------------------------------------------------------- diagnosis
def _resolver(model: WorkbookModel):
    def resolve(sheet: str, row: int, col: int):
        s = model.resolve_sheet(sheet)
        if s is None:
            return "#REF!"
        return model.ws_v(s).cell(row, col).value
    return resolve


def diagnose_cell(model: WorkbookModel, sheet: str, coord: str) -> dict:
    resolved = model.resolve_sheet(sheet)
    if resolved is None:
        raise AppError("sheet_not_found", f"There is no sheet named '{sheet}'.", 404)
    coord = coord.replace("$", "").upper()
    try:
        coordinate_from_string(coord)
    except ValueError as exc:
        raise AppError("invalid_cell", f"'{coord}' is not a valid cell reference.") from exc
    raw, cached = model.cell(resolved, coord)
    cell_obj = model.ws_f(resolved)[coord]
    if cell_obj.data_type != "f":
        return {"kind": "formula_diagnosis", "sheet": resolved, "cell": coord, "has_formula": False,
                "message": f"{resolved}!{coord} does not contain a formula (value: {V.json_safe(raw)!r})."}
    formula = str(getattr(raw, "text", raw))
    error = cached if V.is_error(cached) else None
    findings: list[str] = []
    fix: str | None = None
    if error == "#DIV/0!" or "/" in formula:
        try:
            ast = _Parser(formula).parse()
            for denom in _denominators(ast):
                if denom[0] == "ref":
                    s, r1, c1, r2, c2 = parse_ref(denom[1], resolved)
                    if (r1, c1) == (r2, c2):
                        val = _resolver(model)(s, r1, c1)
                        if V.is_blank(val) or val == 0:
                            findings.append(f"The divisor {denom[1]} is {'empty' if V.is_blank(val) else 'zero'}.")
                            dref = denom[1]
                            fix = f'=IF(OR({dref}=0,{dref}=""),"",{formula[1:]})'
        except Unsupported:
            pass
    if error and not fix:
        fix = f'=IFERROR({formula[1:]},"")'
    if not error:
        preview, reason = evaluate_formula(formula, resolved, _resolver(model))
        if isinstance(preview, str) and V.is_error(preview):
            error = preview
            findings.append(f"The saved result is {V.json_safe(cached)!r}, but re-evaluating with current inputs gives {preview}.")
    explanation = explain_formula(formula, model, resolved, coord)
    return {
        "kind": "formula_diagnosis", "sheet": resolved, "cell": coord, "has_formula": True, "formula": formula,
        "cached_value": V.json_safe(cached), "error": error,
        "meaning": ERROR_HELP.get(error or "", "No error detected in the saved result."),
        "findings": findings, "explanation": explanation, "suggested_formula": fix if error else None,
        "fix_note": ("Wrapping in IF/IFERROR hides the error by returning an empty result. Check the input data "
                     "as well; an empty result may hide a genuine data problem.") if error else None,
    }


def _denominators(node) -> list:
    out = []
    if node[0] == "bin":
        if node[1] == "/":
            out.append(node[3])
        out += _denominators(node[2]) + _denominators(node[3])
    elif node[0] == "fn":
        for a in node[2]:
            out += _denominators(a)
    elif node[0] == "neg":
        out += _denominators(node[1])
    return out


def consistency_check(model: WorkbookModel, sheet: str, column: str) -> dict:
    table = require_table(model, sheet)
    col = require_column(table, column)
    origin = f"{col.letter}{table.first_data_row}"
    patterns: dict[int, str] = {}
    hardcoded: list[int] = []
    for r, vals in table.rows:
        f = table.formulas.get((r, col.index))
        if f is None:
            if not V.is_blank(vals[table.col_pos(col)]):
                hardcoded.append(r)
            continue
        try:
            patterns[r] = Translator(f, origin=f"{col.letter}{r}").translate_formula(origin)
        except Exception:
            patterns[r] = f
    if not patterns:
        return {"kind": "consistency", "sheet": table.sheet, "column": col.name, "formula_rows": 0,
                "message": f"'{col.name}' contains no formulas.", "inconsistent": [], "hardcoded_rows": []}
    common, n = Counter(patterns.values()).most_common(1)[0]
    inconsistent = [{"row": r, "cell": f"{col.letter}{r}", "formula": table.formulas[(r, col.index)],
                     "expected": Translator(common, origin=origin).translate_formula(f"{col.letter}{r}")}
                    for r, p in patterns.items() if p != common]
    return {
        "kind": "consistency", "sheet": table.sheet, "column": col.name, "formula_rows": len(patterns),
        "dominant_pattern": common, "dominant_share": round(n / len(patterns), 3),
        "inconsistent": inconsistent[:50], "hardcoded_rows": hardcoded[:50] if len(patterns) >= len(hardcoded) else [],
        "message": (f"{len(inconsistent)} of {len(patterns)} formulas in '{col.name}' differ from the dominant pattern"
                    + (f"; {len(hardcoded)} row(s) hold typed values instead of formulas." if hardcoded and patterns else ".")),
    }


# --------------------------------------------------------------------------- proposals
def next_free_column(table: SheetTable, model: WorkbookModel) -> int:
    ws = model.ws_f(table.sheet)
    c = table.last_col + 1
    while any(ws.cell(r, c).value not in (None, "") for r in range(table.header_row, table.last_row + 1)):
        c += 1
    return c


def _ref(table: SheetTable, name: str) -> str:
    return require_column(table, name).letter


def build_formula_proposal(model: WorkbookModel, sheet: str, spec: dict, excel_version: str = "365") -> dict:
    """Turn a structured formula spec into a concrete per-row formula column proposal."""
    table = require_table(model, sheet)
    kind = spec.get("kind")
    name = (spec.get("new_column") or "").strip()
    assumptions: list[str] = []
    number_format = "General"
    inputs: list[str] = []
    if kind == "ratio":
        num, den = require_column(table, spec["numerator"]), require_column(table, spec["denominator"])
        template = f'=IFERROR({num.letter}{{r}}/{den.letter}{{r}},"")'
        inputs = [num.name, den.name]
        name = name or f"{num.name} / {den.name}"
        number_format = "0.0%" if spec.get("as_percent", True) else "0.00"
        behavior = f"Each row shows {num.name} divided by {den.name}. Rows where {den.name} is zero or blank show an empty cell instead of #DIV/0!."
        assumptions.append(f"Uses {num.name} and {den.name} as stored in each row (no currency conversion).")
    elif kind == "margin":
        if spec.get("profit"):
            prof, rev = require_column(table, spec["profit"]), require_column(table, spec["revenue"])
            template = f'=IFERROR({prof.letter}{{r}}/{rev.letter}{{r}},"")'
            inputs = [prof.name, rev.name]
            assumptions.append(f"Margin = {prof.name} ÷ {rev.name}.")
        else:
            rev, cost = require_column(table, spec["revenue"]), require_column(table, spec["cost"])
            template = f'=IFERROR(({rev.letter}{{r}}-{cost.letter}{{r}})/{rev.letter}{{r}},"")'
            inputs = [rev.name, cost.name]
            assumptions.append(f"Margin = ({rev.name} − {cost.name}) ÷ {rev.name}; no profit column was found, so profit is derived.")
        name = name or "Profit Margin"
        number_format = "0.0%"
        behavior = "Shows profit as a percentage of revenue for each row; blank when revenue is zero or empty."
    elif kind in ("difference", "sum", "product"):
        a, b = require_column(table, spec["a"]), require_column(table, spec["b"])
        op = {"difference": "-", "sum": "+", "product": "*"}[kind]
        template = f"={a.letter}{{r}}{op}{b.letter}{{r}}"
        inputs = [a.name, b.name]
        name = name or {"difference": f"{a.name} - {b.name}", "sum": f"{a.name} + {b.name}", "product": f"{a.name} x {b.name}"}[kind]
        number_format = a.number_format or "#,##0.00"
        behavior = f"Each row computes {a.name} {op} {b.name}. Blank inputs count as zero; text inputs produce #VALUE!."
    elif kind == "conditional":
        c = require_column(table, spec["column"])
        op = spec.get("op", ">")
        if op not in (">", ">=", "<", "<=", "=", "<>"):
            raise AppError("invalid_formula_spec", f"Unsupported comparison '{op}'.")
        threshold = float(spec["threshold"])
        t_label, f_label = spec.get("true_label", "Yes"), spec.get("false_label", "No")
        thr = int(threshold) if threshold.is_integer() else threshold
        template = f'=IF({c.letter}{{r}}{op}{thr},"{t_label}","{f_label}")'
        inputs = [c.name]
        name = name or f"{c.name} flag"
        behavior = f'Shows "{t_label}" when {c.name} {op} {thr}, otherwise "{f_label}". Blank cells count as zero.'
    elif kind == "date_part":
        c = require_column(table, spec["column"])
        part = spec.get("part", "year_month")
        template = {"year": f"=YEAR({c.letter}{{r}})", "month": f"=MONTH({c.letter}{{r}})",
                    "year_month": f'=TEXT({c.letter}{{r}},"yyyy-mm")'}[part]
        inputs = [c.name]
        name = name or {"year": "Year", "month": "Month", "year_month": "Year-Month"}[part]
        behavior = f"Extracts the {part.replace('_', '-')} from {c.name}. Dates stored as text produce #VALUE!; standardize them first."
    elif kind == "lookup":
        key = require_column(table, spec["key"])
        other = require_table(model, spec["lookup_sheet"])
        okey, oret = require_column(other, spec["lookup_key"]), require_column(other, spec["return_column"])
        sref = other.sheet if re.match(r"^[A-Za-z_][A-Za-z0-9_]*$", other.sheet) else "'" + other.sheet.replace("'", "''") + "'"
        keys = f"{sref}!${okey.letter}${other.first_data_row}:${okey.letter}${other.last_row}"
        rets = f"{sref}!${oret.letter}${other.first_data_row}:${oret.letter}${other.last_row}"
        if excel_version == "365":
            template = f'=XLOOKUP({key.letter}{{r}},{keys},{rets},"Not found")'
        else:
            template = f'=IFERROR(INDEX({rets},MATCH({key.letter}{{r}},{keys},0)),"Not found")'
        inputs = [key.name, f"{other.sheet}!{okey.name}", f"{other.sheet}!{oret.name}"]
        name = name or oret.name
        behavior = f'Finds each row\'s {key.name} in {other.sheet}!{okey.name} and returns {oret.name}; shows "Not found" when there is no match.'
        assumptions.append(f"Exact match on {key.name}; the first match is used if {other.sheet} has duplicates.")
        assumptions.append(f"Lookup range is fixed to rows {other.first_data_row}-{other.last_row}; rows added later to {other.sheet} are not included.")
    elif kind == "custom":
        template = spec.get("template", "")
        if not template.startswith("="):
            raise AppError("invalid_formula_spec", "Custom formulas must start with '='.")

        def sub(m: re.Match) -> str:
            if m.group(1) == "r":
                return "{r}"
            return f"{require_column(table, m.group(1)).letter}{{r}}"

        inputs = [i for i in re.findall(r"\{([^{}]+)\}", template) if i != "r"]
        template = re.sub(r"\{([^{}]+)\}", sub, template)
        name = name or "Calculated"
        behavior = spec.get("behavior") or "Computes the custom expression for each row."
        assumptions.append("Custom formula supplied by the request; review it carefully.")
    else:
        raise AppError("invalid_formula_spec", f"Unknown formula kind '{kind}'.")

    if table.column(name):
        raise AppError("column_exists", f"A column named '{name}' already exists on {table.sheet}. Choose another name.", 409)
    target_col = next_free_column(table, model)
    letter = get_column_letter(target_col)
    data_rows = [r for r, _ in table.rows]
    first, last = (data_rows[0], data_rows[-1]) if data_rows else (table.first_data_row, table.first_data_row)
    sample_formula = template.replace("{r}", str(first))
    notes = validate_formula(sample_formula, model)
    explanation = explain_formula(sample_formula, model, table.sheet, f"{letter}{first}")
    return {
        "kind": "formula_proposal",
        "sheet": table.sheet,
        "new_column": name,
        "target_column": letter,
        "target_range": f"{letter}{first}:{letter}{last}",
        "header_cell": f"{letter}{table.header_row}",
        "template": template,
        "first_formula": sample_formula,
        "rows": data_rows,
        "number_format": number_format,
        "inputs": inputs,
        "explanation": explanation["plain_english"],
        "functions": explanation["functions"],
        "references": explanation["references"],
        "assumptions": assumptions,
        "expected_behavior": behavior,
        "compatibility": compatibility(sample_formula) + notes,
        "anchoring": "Row references are relative (they shift for each row); lookup ranges are absolute ($) so they stay fixed.",
    }


def column_index(letter: str) -> int:
    return column_index_from_string(letter)
