"""Deterministic analytics over detected sheet tables.

Every number FloPilot reports is computed here. Each result carries its calculation method,
filters, missing-data treatment, exclusions and source references so the answer can be audited.
"""
from __future__ import annotations

import difflib
import math
import re
import statistics
from collections import defaultdict
from datetime import date, datetime
from typing import Any, Iterable

from ..common import AppError
from . import values as V
from .model import Column, SheetTable, WorkbookModel
from .sources import make_source

AGG_FUNCS = {"sum", "average", "count", "min", "max", "median", "distinct_count"}

SYNONYMS = {
    "revenue": ["revenue", "sales", "turnover", "gross revenue", "net revenue", "income"],
    "cost": ["cost", "cogs", "expense", "spend", "amount"],
    "profit": ["profit", "margin", "earnings"],
    "date": ["date", "order date", "invoice date", "month", "period"],
    "quantity": ["units", "quantity", "qty", "volume"],
}


# --------------------------------------------------------------------------- column matching
def match_columns(table: SheetTable, phrase: str) -> list[Column]:
    """Columns that plausibly correspond to `phrase` (best first). Several results = ambiguity."""
    p = phrase.strip().lower()
    if not p:
        return []
    exact = [c for c in table.columns if c.name.lower() == p]
    if exact:
        return exact
    contains = [c for c in table.columns if re.search(rf"\b{re.escape(p)}\b", c.name.lower())]
    if contains:
        return contains
    singular = p[:-1] if p.endswith("s") else p
    contains = [c for c in table.columns if re.search(rf"\b{re.escape(singular)}", c.name.lower())]
    if contains:
        return contains
    for key, words in SYNONYMS.items():
        if p == key or p in words:
            hits = [c for c in table.columns if any(w in c.name.lower() for w in words)]
            if hits:
                return hits
    names = [c.name.lower() for c in table.columns]
    close = difflib.get_close_matches(p, names, n=3, cutoff=0.75)
    return [c for c in table.columns if c.name.lower() in close]


def require_column(table: SheetTable, name: str) -> Column:
    col = table.column(name)
    if col is None:
        matches = match_columns(table, name)
        if len(matches) == 1:
            return matches[0]
        available = ", ".join(c.name for c in table.columns)
        if matches:
            raise AppError("ambiguous_column", f"'{name}' could refer to several columns on {table.sheet}: "
                           f"{', '.join(c.name for c in matches)}.", 422, {"candidates": [c.name for c in matches]})
        raise AppError("column_not_found", f"There is no column '{name}' on sheet {table.sheet}. Available columns: {available}.", 404)
    return col


def require_table(model: WorkbookModel, sheet: str) -> SheetTable:
    resolved = model.resolve_sheet(sheet)
    if resolved is None:
        raise AppError("sheet_not_found", f"There is no sheet named '{sheet}'. Sheets: {', '.join(model.sheet_names)}.", 404)
    table = model.table(resolved)
    if table is None:
        raise AppError("empty_sheet", f"Sheet '{resolved}' has no tabular data.", 422)
    return table


# --------------------------------------------------------------------------- filters
FILTER_OPS = {"eq", "neq", "gt", "gte", "lt", "lte", "contains", "in", "between", "not_blank", "is_blank"}


def _cmp_value(v: Any) -> Any:
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, str):
        return v.strip().lower()
    return v


def _coerce_like(target: Any, sample: Any) -> Any:
    if isinstance(sample, (date, datetime)) and isinstance(target, str):
        parsed, _ = V.parse_date_text(target, None)
        return parsed or target
    if V.is_number(sample) and isinstance(target, str):
        number = V.parse_numeric_text(target)
        return number if number is not None else target
    return target


def _match(value: Any, op: str, target: Any) -> bool:
    if op == "is_blank":
        return V.is_blank(value)
    if op == "not_blank":
        return not V.is_blank(value)
    if V.is_blank(value):
        return False
    v = _cmp_value(value)
    if op == "in":
        return v in {_cmp_value(_coerce_like(t, value)) for t in target}
    if op == "between":
        lo, hi = (_cmp_value(_coerce_like(t, value)) for t in target)
        try:
            return lo <= v <= hi
        except TypeError:
            return False
    t = _cmp_value(_coerce_like(target, value))
    try:
        if op == "eq":
            return v == t
        if op == "neq":
            return v != t
        if op == "contains":
            return str(t) in str(v)
        if op == "gt":
            return v > t
        if op == "gte":
            return v >= t
        if op == "lt":
            return v < t
        if op == "lte":
            return v <= t
    except TypeError:
        return False
    raise AppError("invalid_filter", f"Unsupported filter operator '{op}'.")


def apply_filters(table: SheetTable, filters: Iterable[dict] | None) -> tuple[list[tuple[int, list[Any]]], list[dict]]:
    rows = table.rows
    described = []
    for f in filters or []:
        op = f.get("op", "eq")
        if op not in FILTER_OPS:
            raise AppError("invalid_filter", f"Unsupported filter operator '{op}'.")
        col = require_column(table, f["column"])
        pos = table.col_pos(col)
        rows = [(r, vals) for r, vals in rows if _match(vals[pos], op, f.get("value"))]
        described.append({"column": col.name, "op": op, "value": f.get("value")})
    return rows, described


# --------------------------------------------------------------------------- aggregation
def _aggregate(func: str, nums: list[float], raw_non_blank: int) -> float | int | None:
    if func == "count":
        return raw_non_blank
    if not nums:
        return None
    if func == "sum":
        return math.fsum(nums)
    if func == "average":
        return math.fsum(nums) / len(nums)
    if func == "min":
        return min(nums)
    if func == "max":
        return max(nums)
    if func == "median":
        return statistics.median(nums)
    raise AppError("invalid_aggregation", f"Unsupported aggregation '{func}'.")


def _numeric_split(values: list[Any]) -> tuple[list[float], dict]:
    nums: list[float] = []
    excluded = {"blank": 0, "text": 0, "numeric_text": 0, "error": 0}
    for v in values:
        kind = V.classify(v)
        if kind == "number":
            nums.append(float(v))
        elif kind == "blank":
            excluded["blank"] += 1
        elif kind == "error":
            excluded["error"] += 1
        elif kind == "numeric_text":
            excluded["numeric_text"] += 1
        else:
            excluded["text"] += 1
    return nums, excluded


def aggregate(model: WorkbookModel, sheet: str, measure: str | None, func: str = "sum",
              group_by: list[str] | None = None, filters: list[dict] | None = None,
              sort: str = "desc", limit: int = 50) -> dict:
    func = func.lower()
    if func not in AGG_FUNCS:
        raise AppError("invalid_aggregation", f"Unsupported aggregation '{func}'. Use one of {sorted(AGG_FUNCS)}.")
    table = require_table(model, sheet)
    rows, described = apply_filters(table, filters)
    measure_col = require_column(table, measure) if measure else None
    group_cols = [require_column(table, g) for g in group_by or []]
    if measure_col is None and func != "count":
        raise AppError("measure_required", f"A measure column is required for '{func}'.")

    groups: dict[tuple, list[tuple[int, list[Any]]]] = defaultdict(list)
    display_key: dict[tuple, tuple] = {}
    for r, vals in rows:
        raw_key = tuple(vals[table.col_pos(g)] for g in group_cols)
        key = tuple(_cmp_value(k) if not V.is_blank(k) else None for k in raw_key)
        groups[key].append((r, vals))
        display_key.setdefault(key, tuple(V.json_safe(k.strip() if isinstance(k, str) else k) if not V.is_blank(k) else "(blank)" for k in raw_key))

    total_excluded = {"blank": 0, "text": 0, "numeric_text": 0, "error": 0}
    out_rows = []
    for key, members in groups.items():
        if measure_col is not None:
            pos = table.col_pos(measure_col)
            vals = [m[1][pos] for m in members]
            nums, excluded = _numeric_split(vals)
            for k, n in excluded.items():
                total_excluded[k] += n
            raw_non_blank = sum(not V.is_blank(v) for v in vals) if func == "count" else len(nums)
            if func == "distinct_count":
                value: Any = len({_cmp_value(v) for v in vals if not V.is_blank(v)})
            else:
                value = _aggregate(func, nums, raw_non_blank)
        else:
            value = len(members)
        out_rows.append({"group": list(display_key[key]), "value": _round(value), "rows": len(members)})

    reverse = sort != "asc"
    out_rows.sort(key=lambda x: (x["value"] is None, -(x["value"] or 0) if reverse else (x["value"] or 0)))
    truncated = len(out_rows) > limit
    grand_nums = [] if measure_col is None else _numeric_split([v[table.col_pos(measure_col)] for _, v in rows])[0]
    total = _round(_aggregate(func, grand_nums, len(grand_nums))) if measure_col is not None and func in ("sum", "count") else None

    sources = [make_source(model, table.sheet, table.range_ref(measure_col), measure_col.name)] if measure_col else []
    sources += [make_source(model, table.sheet, table.range_ref(g), g.name) for g in group_cols]
    sources += [make_source(model, table.sheet, table.range_ref(require_column(table, f["column"])), f["column"]) for f in described]
    if not sources:
        sources = [make_source(model, table.sheet, table.range_ref())]

    warnings = []
    if total_excluded["numeric_text"]:
        warnings.append(f"{total_excluded['numeric_text']} value(s) stored as text that look like numbers were excluded. Convert them to numbers to include them.")
    if total_excluded["text"]:
        warnings.append(f"{total_excluded['text']} non-numeric text value(s) were excluded.")
    if total_excluded["error"]:
        warnings.append(f"{total_excluded['error']} error value(s) (e.g. #DIV/0!) were excluded.")
    if measure_col is not None:
        units = _unit_symbols(model, table, measure_col, [r for r, _ in rows])
        if len(units) > 1:
            warnings.append(f"'{measure_col.name}' mixes currency formats ({', '.join(sorted(units))}); values were added as-is "
                            "without conversion. Check that they are in the same unit.")
    if group_cols and any(g.semantic == "category" for g in group_cols):
        from .quality import category_variants
        for g in group_cols:
            if category_variants(table, g):
                warnings.append(f"'{g.name}' contains spelling variants that are grouped separately. Normalize them first for combined totals.")
    if any("(blank)" in r["group"] for r in out_rows):
        warnings.append("Rows with a blank grouping value are shown as '(blank)'.")

    return {
        "kind": "aggregate",
        "sheet": table.sheet,
        "measure": measure_col.name if measure_col else None,
        "function": func,
        "group_by": [g.name for g in group_cols],
        "rows": out_rows[:limit],
        "truncated": truncated,
        "group_count": len(out_rows),
        "total": total,
        "rows_considered": len(rows),
        "filters": described,
        "excluded": total_excluded,
        "method": _method_text(func, measure_col, group_cols, table),
        "missing_treatment": "Blank cells are excluded from sums and averages; text and error values in the measure column are excluded and counted.",
        "sources": sources,
        "warnings": warnings,
        "number_format": measure_col.number_format if measure_col else None,
        "semantic": measure_col.semantic if measure_col else None,
    }


def _unit_symbols(model: WorkbookModel, table: SheetTable, col: Column, row_numbers: list[int]) -> set[str]:
    ws = model.ws_f(table.sheet)
    found: set[str] = set()
    for r in row_numbers[:20000]:
        fmt = ws.cell(r, col.index).number_format or ""
        for sym in ("$", "€", "£", "¥"):
            if sym in fmt:
                found.add(sym)
    return found


def _method_text(func: str, measure: Column | None, groups: list[Column], table: SheetTable) -> str:
    target = f"'{measure.name}' ({table.range_ref(measure)})" if measure else "rows"
    by = f" grouped by {', '.join(repr(g.name) for g in groups)}" if groups else ""
    return f"{func.upper()} of {target} on sheet '{table.sheet}'{by}, computed by FloPilot's deterministic engine from cell values (cached results for formula cells)."


def _round(value: Any) -> Any:
    if isinstance(value, float):
        return round(value, 6)
    return value


def missing_by_column(model: WorkbookModel, sheet: str) -> dict:
    table = require_table(model, sheet)
    rows = []
    for col in table.columns:
        vals = table.values(col)
        missing = sum(V.is_blank(v) for v in vals)
        rows.append({"group": [col.name], "value": missing, "pct": round(100 * missing / len(vals), 1) if vals else 0.0, "rows": len(vals)})
    rows.sort(key=lambda r: -r["value"])
    return {
        "kind": "missing", "sheet": table.sheet, "rows": rows, "rows_considered": len(table.rows),
        "method": f"Count of blank cells (empty or whitespace-only) per column in {table.sheet}!{table.range_ref()}.",
        "missing_treatment": "Formula cells with a saved empty-string result count as blank.",
        "filters": [], "sources": [make_source(model, table.sheet, table.range_ref())], "warnings": [],
    }


def _month(d: date) -> str:
    return f"{d.year:04d}-{d.month:02d}"


def _date_of(value: Any) -> date | None:
    if isinstance(value, (date, datetime)):
        return V.as_date(value)
    return None


def time_series(model: WorkbookModel, sheet: str, date_column: str, measure: str, func: str = "sum",
                filters: list[dict] | None = None, group_by: str | None = None) -> dict:
    table = require_table(model, sheet)
    date_col = require_column(table, date_column)
    measure_col = require_column(table, measure)
    group_col = require_column(table, group_by) if group_by else None
    rows, described = apply_filters(table, filters)
    buckets: dict[tuple[str, str], list[float]] = defaultdict(list)
    undated = 0
    text_dates = 0
    excluded_total = {"blank": 0, "text": 0, "numeric_text": 0, "error": 0}
    for _, vals in rows:
        d = _date_of(vals[table.col_pos(date_col)])
        if d is None:
            undated += 1
            if isinstance(vals[table.col_pos(date_col)], str) and V.date_text_kind(vals[table.col_pos(date_col)]):
                text_dates += 1
            continue
        nums, excluded = _numeric_split([vals[table.col_pos(measure_col)]])
        for k, n in excluded.items():
            excluded_total[k] += n
        g = str(vals[table.col_pos(group_col)]).strip() if group_col else "All"
        if nums:
            buckets[(g, _month(d))].extend(nums)
    months = sorted({m for _, m in buckets})
    all_months = _month_range(months[0], months[-1]) if months else []
    series: dict[str, list[dict]] = defaultdict(list)
    for g in sorted({g for g, _ in buckets}):
        for m in all_months:
            nums = buckets.get((g, m))
            series[g].append({"period": m, "value": _round(_aggregate(func, nums, len(nums))) if nums else None})
    gaps = [m for m in all_months if not any((g, m) in buckets for g in series)]
    warnings = []
    if undated:
        warnings.append(f"{undated} row(s) without a usable date were excluded" + (f" ({text_dates} have dates stored as text; standardize them to include)." if text_dates else "."))
    if gaps:
        warnings.append(f"No data for {len(gaps)} month(s): {', '.join(gaps[:6])}{'…' if len(gaps) > 6 else ''}. Shown as gaps, not zeros.")
    return {
        "kind": "time_series", "sheet": table.sheet, "measure": measure_col.name, "function": func,
        "date_column": date_col.name, "group_by": [group_col.name] if group_col else [],
        "periods": all_months, "series": dict(series), "rows_considered": len(rows), "filters": described,
        "excluded": excluded_total, "undated_rows": undated,
        "method": f"{func.upper()} of '{measure_col.name}' per calendar month of '{date_col.name}' on '{table.sheet}'.",
        "missing_treatment": "Rows without a real date value are excluded; months with no rows are gaps (null), not zero.",
        "sources": [make_source(model, table.sheet, table.range_ref(date_col), date_col.name),
                    make_source(model, table.sheet, table.range_ref(measure_col), measure_col.name)],
        "warnings": warnings, "number_format": measure_col.number_format, "semantic": measure_col.semantic,
    }


def _month_range(start: str, end: str) -> list[str]:
    y, m = map(int, start.split("-"))
    ey, em = map(int, end.split("-"))
    out = []
    while (y, m) <= (ey, em):
        out.append(f"{y:04d}-{m:02d}")
        m += 1
        if m > 12:
            y, m = y + 1, 1
    return out


def period_change(model: WorkbookModel, sheet: str, date_column: str, measure: str, group_by: str,
                  limit: int = 10) -> dict:
    """Month-over-month change for the two most recent months that have data."""
    ts = time_series(model, sheet, date_column, measure, "sum", group_by=group_by)
    periods = [p for p in ts["periods"] if any(
        next((pt["value"] for pt in s if pt["period"] == p), None) is not None for s in ts["series"].values())]
    if len(periods) < 2:
        raise AppError("insufficient_periods", "At least two months with data are needed for a month-over-month comparison.", 422)
    prev, last = periods[-2], periods[-1]
    rows = []
    for g, pts in ts["series"].items():
        lookup = {pt["period"]: pt["value"] for pt in pts}
        a, b = lookup.get(prev), lookup.get(last)
        change = None if a is None or b is None else _round(b - a)
        pct = None if a in (None, 0) or b is None else round(100 * (b - a) / abs(a), 2)
        rows.append({"group": [g], "previous": a, "current": b, "value": change, "pct": pct})
    rows.sort(key=lambda r: (r["value"] is None, r["value"] if r["value"] is not None else 0))
    ts.update({
        "kind": "period_change", "previous_period": prev, "current_period": last, "rows": rows[:limit],
        "method": f"SUM of '{ts['measure']}' by '{group_by}' for {prev} vs {last} (latest two months with data); change = current - previous.",
    })
    if any(r["value"] is None for r in rows):
        ts["warnings"].append("Groups missing data in either month have no change value.")
    return ts


def compare_measures(model: WorkbookModel, sheet: str, actual: str, budget: str, group_by: str | None) -> dict:
    table = require_table(model, sheet)
    a_col, b_col = require_column(table, actual), require_column(table, budget)
    g_col = require_column(table, group_by) if group_by else None
    acc: dict[str, list[float]] = defaultdict(lambda: [0.0, 0.0, 0])
    skipped = 0
    for _, vals in table.rows:
        a, b = vals[table.col_pos(a_col)], vals[table.col_pos(b_col)]
        if not (V.is_number(a) and V.is_number(b)):
            skipped += 1
            continue
        key = str(vals[table.col_pos(g_col)]).strip() if g_col else "All"
        acc[key][0] += float(a)
        acc[key][1] += float(b)
        acc[key][2] += 1
    rows = []
    for key, (a, b, n) in acc.items():
        rows.append({"group": [key], "actual": _round(a), "budget": _round(b), "value": _round(a - b),
                     "pct": round(100 * (a - b) / b, 2) if b else None, "rows": n})
    rows.sort(key=lambda r: -(r["value"] or 0))
    warnings = [f"{skipped} row(s) without numeric values in both columns were excluded."] if skipped else []
    return {
        "kind": "variance", "sheet": table.sheet, "actual": a_col.name, "budget": b_col.name,
        "group_by": [g_col.name] if g_col else [], "rows": rows, "rows_considered": len(table.rows) - skipped,
        "total": {"actual": _round(sum(r["actual"] for r in rows)), "budget": _round(sum(r["budget"] for r in rows))},
        "method": f"Variance = SUM('{a_col.name}') - SUM('{b_col.name}')" + (f" per '{g_col.name}'" if g_col else "") + "; % = variance / budget.",
        "missing_treatment": "Rows lacking a number in either column are excluded.", "filters": [],
        "sources": [make_source(model, table.sheet, table.range_ref(a_col), a_col.name),
                    make_source(model, table.sheet, table.range_ref(b_col), b_col.name)],
        "warnings": warnings, "number_format": a_col.number_format, "semantic": a_col.semantic,
    }


def correlation(model: WorkbookModel, sheet: str, x: str, y: str) -> dict:
    table = require_table(model, sheet)
    xc, yc = require_column(table, x), require_column(table, y)
    pairs = [(float(v[table.col_pos(xc)]), float(v[table.col_pos(yc)])) for _, v in table.rows
             if V.is_number(v[table.col_pos(xc)]) and V.is_number(v[table.col_pos(yc)])]
    if len(pairs) < 3:
        raise AppError("insufficient_data", "At least three rows with numbers in both columns are needed.", 422)
    r = statistics.correlation([p[0] for p in pairs], [p[1] for p in pairs])
    return {
        "kind": "correlation", "sheet": table.sheet, "x": xc.name, "y": yc.name, "value": round(r, 4), "n": len(pairs),
        "points": [{"x": a, "y": b} for a, b in pairs[:500]],
        "method": "Pearson correlation coefficient over rows with numbers in both columns.",
        "missing_treatment": "Rows missing either value are excluded.", "filters": [], "rows_considered": len(pairs),
        "sources": [make_source(model, table.sheet, table.range_ref(xc), xc.name), make_source(model, table.sheet, table.range_ref(yc), yc.name)],
        "warnings": ["Correlation describes association, not causation: it does not show that one causes the other."],
    }


def histogram(model: WorkbookModel, sheet: str, column: str, bins: int = 10) -> dict:
    table = require_table(model, sheet)
    col = require_column(table, column)
    nums, excluded = _numeric_split(table.values(col))
    if not nums:
        raise AppError("insufficient_data", f"'{col.name}' has no numeric values.", 422)
    lo, hi = min(nums), max(nums)
    width = (hi - lo) / bins or 1
    counts = [0] * bins
    for v in nums:
        counts[min(int((v - lo) / width), bins - 1)] += 1
    rows = [{"group": [f"{lo + i * width:,.0f}–{lo + (i + 1) * width:,.0f}"], "value": c} for i, c in enumerate(counts)]
    return {
        "kind": "histogram", "sheet": table.sheet, "measure": col.name, "rows": rows, "rows_considered": len(nums),
        "method": f"{bins} equal-width bins between min {lo:,.2f} and max {hi:,.2f}.", "excluded": excluded,
        "missing_treatment": "Non-numeric and blank cells are excluded.", "filters": [],
        "sources": [make_source(model, table.sheet, table.range_ref(col), col.name)], "warnings": [],
    }


def read_range(model: WorkbookModel, sheet: str, range_ref: str, max_cells: int = 400) -> dict:
    from openpyxl.utils.cell import range_boundaries

    resolved = model.resolve_sheet(sheet)
    if resolved is None:
        raise AppError("sheet_not_found", f"There is no sheet named '{sheet}'.", 404)
    try:
        min_col, min_row, max_col, max_row = range_boundaries(range_ref.replace("$", ""))
    except (ValueError, TypeError) as exc:
        raise AppError("invalid_range", f"'{range_ref}' is not a valid range.") from exc
    if (max_row - min_row + 1) * (max_col - min_col + 1) > max_cells:
        raise AppError("range_too_large", f"Ranges are limited to {max_cells} cells per read.", 422)
    ws_f, ws_v = model.ws_f(resolved), model.ws_v(resolved)
    cells = []
    for r in range(min_row, max_row + 1):
        row = []
        for c in range(min_col, max_col + 1):
            cf = ws_f.cell(r, c)
            row.append({"value": V.json_safe(ws_v.cell(r, c).value),
                        "formula": str(getattr(cf.value, "text", cf.value)) if cf.data_type == "f" else None})
        cells.append(row)
    return {"kind": "range", "sheet": resolved, "range": range_ref, "cells": cells,
            "sources": [make_source(model, resolved, range_ref)]}
