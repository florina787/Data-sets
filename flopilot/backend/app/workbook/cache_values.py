"""Cached formula results for workbooks written by openpyxl.

openpyxl saves formulas WITHOUT their last calculated values, so a freshly written version would
show every formula as blank to any reader that does not recalculate (including FloPilot's own
analysis). After each edit FloPilot restores cached values:

1. Formulas within the evaluator's supported subset are recomputed from the new inputs.
2. Otherwise, if the formula text is unchanged and none of its direct inputs changed, Excel's
   previously saved value is carried over.
3. Anything else is left uncalculated and reported; Excel recalculates on open
   (fullCalcOnLoad is set on every edited version).
"""
from __future__ import annotations

import io
import re
import zipfile
from typing import Any
from xml.sax.saxutils import escape

import openpyxl

from .evaluator import Evaluator, Unsupported, XLError, parse_ref, referenced_ranges


def _sheet_parts(data: bytes) -> dict[str, str]:
    """Map sheet title -> worksheet part name using the workbook relationships."""
    with zipfile.ZipFile(io.BytesIO(data)) as zf:
        wb_xml = zf.read("xl/workbook.xml").decode("utf-8")
        rels = zf.read("xl/_rels/workbook.xml.rels").decode("utf-8")
    targets = {m.group(1): m.group(2) for m in re.finditer(r'<Relationship[^>]*Id="([^"]+)"[^>]*Target="([^"]+)"', rels)}
    targets.update({m.group(2): m.group(1) for m in re.finditer(r'<Relationship[^>]*Target="([^"]+)"[^>]*Id="([^"]+)"', rels)})
    parts = {}
    for m in re.finditer(r'<sheet [^>]*name="([^"]+)"[^>]*r:id="([^"]+)"', wb_xml):
        name = m.group(1).replace("&amp;", "&").replace("&lt;", "<").replace("&gt;", ">").replace("&quot;", '"').replace("&apos;", "'")
        target = targets.get(m.group(2), "")
        target = target.lstrip("/")
        parts[name] = target if target.startswith("xl/") else f"xl/{target}"
    return parts


def inject_cached_values(data: bytes, values: dict[str, dict[str, Any]]) -> bytes:
    """Write <v> results into formula cells of an openpyxl-written package."""
    parts = _sheet_parts(data)
    by_part = {parts[s]: v for s, v in values.items() if s in parts and v}
    src = zipfile.ZipFile(io.BytesIO(data))
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as dst:
        for item in src.infolist():
            content = src.read(item.filename)
            cells = by_part.get(item.filename)
            if cells:
                text = content.decode("utf-8")

                def repl(m: re.Match, cells: dict[str, Any] = cells) -> str:
                    coord = m.group(1)
                    if coord not in cells:
                        return m.group(0)
                    value = cells[coord]
                    attrs = re.sub(r'\s+t="[^"]*"', "", m.group(2))
                    if isinstance(value, XLError):
                        t, v = ' t="e"', value.code
                    elif isinstance(value, str) and value.strip() in ("#DIV/0!", "#N/A", "#NAME?", "#NULL!", "#NUM!", "#REF!", "#VALUE!"):
                        t, v = ' t="e"', value.strip()
                    elif isinstance(value, bool):
                        t, v = ' t="b"', "1" if value else "0"
                    elif isinstance(value, str):
                        t, v = ' t="str"', escape(value)
                    elif isinstance(value, (int, float)):
                        t, v = "", repr(float(value))
                    else:
                        return m.group(0)
                    return f'<c r="{coord}"{attrs}{t}><f>{m.group(3)}</f><v>{v}</v></c>'

                text = re.sub(r'<c r="([A-Z]{1,3}\d+)"([^>]*)><f>(.*?)</f><v\s*/?>(?:</v>)?</c>', repl, text)
                content = text.encode("utf-8")
            dst.writestr(item, content)
    return out.getvalue()


def refresh_cached_values(base: bytes, candidate: bytes, change_set: dict) -> tuple[bytes, dict]:
    wb_new = openpyxl.load_workbook(io.BytesIO(candidate))
    wb_old_f = openpyxl.load_workbook(io.BytesIO(base))
    wb_old_v = openpyxl.load_workbook(io.BytesIO(base), data_only=True)
    changed = {(c["sheet"], c["cell"]) for c in change_set.get("changes", [])}
    structural_sheet = change_set.get("sheet") if (change_set.get("structural") or {}).get("delete_rows") else None
    memo: dict[tuple[str, int, int], Any] = {}
    how: dict[tuple[str, str], str] = {}
    stack: set[tuple[str, int, int]] = set()

    def raw(sheet: str, row: int, col: int):
        if sheet not in wb_new.sheetnames:
            return "#REF!"
        return wb_new[sheet].cell(row=row, column=col)

    def resolve(sheet: str, row: int, col: int):
        key = (sheet, row, col)
        if key in memo:
            return memo[key]
        cell = raw(sheet, row, col)
        if isinstance(cell, str):
            return cell
        if cell.data_type != "f":
            return cell.value
        if key in stack:
            raise Unsupported("circular reference")
        stack.add(key)
        try:
            value = compute(sheet, cell)
        finally:
            stack.discard(key)
        memo[key] = value
        return value

    def carried(sheet: str, cell) -> Any:
        coord = cell.coordinate
        if sheet == structural_sheet or (sheet, coord) in changed or sheet not in wb_old_f.sheetnames:
            raise Unsupported("no reusable cached value")
        old = wb_old_f[sheet][coord]
        if old.data_type != "f" or str(getattr(old.value, "text", old.value)) != str(getattr(cell.value, "text", cell.value)):
            raise Unsupported("formula changed")
        for ref in referenced_ranges(str(cell.value)):
            s, r1, c1, r2, c2 = parse_ref(ref, sheet)
            if s == structural_sheet:
                raise Unsupported("inputs moved")
            for (cs, cc) in changed:
                if cs != s:
                    continue
                from openpyxl.utils.cell import coordinate_from_string, column_index_from_string
                col_l, row_n = coordinate_from_string(cc)
                if r1 <= row_n <= r2 and c1 <= column_index_from_string(col_l) <= c2:
                    raise Unsupported("inputs changed")
        value = wb_old_v[sheet][coord].value
        if value is None:
            raise Unsupported("no saved value")
        how[(sheet, coord)] = "carried_over"
        return value

    def compute(sheet: str, cell) -> Any:
        formula = str(getattr(cell.value, "text", cell.value))
        try:
            value = Evaluator(resolve, sheet).evaluate(formula)
            how[(sheet, cell.coordinate)] = "recomputed"
            return value
        except Unsupported:
            return carried(sheet, cell)

    results: dict[str, dict[str, Any]] = {}
    stats = {"recomputed": 0, "carried_over": 0, "uncalculated": 0}
    for ws in wb_new.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                if cell.data_type != "f" or not isinstance(cell.value, str):
                    continue
                try:
                    value = resolve(ws.title, cell.row, cell.column)
                except Exception:  # Unsupported, circular references, unexpected shapes
                    stats["uncalculated"] += 1
                    continue
                if isinstance(value, list) or value is None:
                    stats["uncalculated"] += 1
                    continue
                results.setdefault(ws.title, {})[cell.coordinate] = value
    for sheet, cells in results.items():
        for coord in cells:
            stats[how.get((sheet, coord), "recomputed")] += 1
    return inject_cached_values(candidate, results), stats
