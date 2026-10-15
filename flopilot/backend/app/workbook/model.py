"""In-memory workbook model with header/table detection.

Each stored version is loaded twice with openpyxl: once for formulas (the expressions) and once
with `data_only=True` (the *cached* results Excel last saved). The two are kept separate so we
never confuse an expression with its value. Loaded models are cached per (workbook, version);
because versions are immutable, the cache never needs invalidation, only eviction.
"""
from __future__ import annotations

import csv
import io
import threading
from collections import OrderedDict
from dataclasses import dataclass, field
from typing import Any

import openpyxl
from openpyxl.utils import get_column_letter
from openpyxl.workbook.workbook import Workbook

from ..common import AppError
from ..config import get_settings
from . import values as V
from .ingest import decode_csv


@dataclass
class Column:
    index: int  # 1-based column index in the sheet
    letter: str
    name: str
    inferred_type: str = "unknown"
    semantic: str | None = None  # currency | percentage | date | identifier | category | None
    type_confidence: float = 1.0
    number_format: str | None = None

    def ref(self) -> str:
        return self.letter


@dataclass
class SheetTable:
    """A rectangular, header-led data region on one sheet."""

    sheet: str
    header_row: int
    first_data_row: int
    last_row: int
    first_col: int
    last_col: int
    columns: list[Column]
    rows: list[tuple[int, list[Any]]]  # (row_number, values aligned with columns)
    formulas: dict[tuple[int, int], str] = field(default_factory=dict)  # (row, col) -> "=..."
    header_confidence: float = 1.0
    header_issues: list[str] = field(default_factory=list)
    excel_table: str | None = None

    def column(self, name: str) -> Column | None:
        lowered = name.strip().lower()
        for col in self.columns:
            if col.name.strip().lower() == lowered:
                return col
        return None

    def col_pos(self, col: Column) -> int:
        return self.columns.index(col)

    def values(self, col: Column) -> list[Any]:
        pos = self.col_pos(col)
        return [vals[pos] for _, vals in self.rows]

    def range_ref(self, col: Column | None = None) -> str:
        if col is None:
            return f"{get_column_letter(self.first_col)}{self.header_row}:{get_column_letter(self.last_col)}{self.last_row}"
        return f"{col.letter}{self.first_data_row}:{col.letter}{self.last_row}"


class WorkbookModel:
    def __init__(self, wb_formulas: Workbook, wb_values: Workbook, workbook_id: str = "", version: int = 0):
        self.wb_f = wb_formulas
        self.wb_v = wb_values
        self.workbook_id = workbook_id
        self.version = version
        self._tables: dict[str, SheetTable | None] = {}
        self._lock = threading.Lock()

    @property
    def sheet_names(self) -> list[str]:
        return list(self.wb_f.sheetnames)

    def ws_f(self, sheet: str):
        if sheet not in self.wb_f.sheetnames:
            raise AppError("sheet_not_found", f"There is no sheet named '{sheet}'. Sheets: {', '.join(self.sheet_names)}.", 404)
        return self.wb_f[sheet]

    def ws_v(self, sheet: str):
        self.ws_f(sheet)
        return self.wb_v[sheet]

    def resolve_sheet(self, name: str | None) -> str | None:
        if not name:
            return None
        for sheet in self.sheet_names:
            if sheet.lower() == name.strip().lower():
                return sheet
        return None

    def table(self, sheet: str) -> SheetTable | None:
        with self._lock:
            if sheet not in self._tables:
                self._tables[sheet] = detect_table(self, sheet)
            return self._tables[sheet]

    def cell(self, sheet: str, coord: str) -> tuple[Any, Any]:
        """(formula-or-value, cached value) for one cell."""
        raw = self.ws_f(sheet)[coord].value
        cached = self.ws_v(sheet)[coord].value
        return raw, cached


def formula_text(raw: Any) -> str:
    """Formula expression for a formula cell (plain or ArrayFormula)."""
    text = getattr(raw, "text", None)
    return str(text if text is not None else raw)


def load_model(data: bytes, workbook_id: str = "", version: int = 0) -> WorkbookModel:
    try:
        wb_f = openpyxl.load_workbook(io.BytesIO(data), data_only=False, keep_links=True)
        wb_v = openpyxl.load_workbook(io.BytesIO(data), data_only=True, keep_links=True)
    except AppError:
        raise
    except Exception as exc:  # openpyxl raises many types for malformed packages
        raise AppError("malformed_file", "The workbook could not be read. It may be damaged or use an unsupported feature.") from exc
    settings = get_settings()
    if len(wb_f.sheetnames) > settings.max_sheets:
        raise AppError("too_many_sheets", f"The workbook has more than {settings.max_sheets} sheets.", 413)
    total = sum((ws.max_row or 0) * (ws.max_column or 0) for ws in wb_f.worksheets)
    if total > settings.max_cells:
        raise AppError("too_many_cells", f"The workbook's used ranges exceed the {settings.max_cells:,}-cell limit for this version of FloPilot.", 413)
    return WorkbookModel(wb_f, wb_v, workbook_id, version)


def csv_to_xlsx(data: bytes, sheet_name: str = "Data") -> bytes:
    """Convert CSV to a single-sheet workbook. Plain numbers become numbers; everything else
    stays text. Text that *looks* like a formula is stored as text, never as a formula."""
    text = decode_csv(data)
    try:
        dialect = csv.Sniffer().sniff(text[:20_000], delimiters=",;|\t")
    except csv.Error:
        dialect = csv.excel
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = sheet_name
    for r, row in enumerate(csv.reader(io.StringIO(text), dialect), start=1):
        for c, raw in enumerate(row, start=1):
            if raw == "":
                continue
            number = V.parse_plain_number(raw) if r > 1 else None
            cell = ws.cell(row=r, column=c)
            if number is not None:
                cell.value = number
            else:
                cell.value = raw
                cell.data_type = "s"  # neutralise formula-looking text from untrusted CSV
    if ws.max_row >= 1:
        for cell in ws[1]:
            cell.font = openpyxl.styles.Font(bold=True)
        ws.freeze_panes = "A2"
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def detect_table(model: WorkbookModel, sheet: str) -> SheetTable | None:
    ws_f = model.ws_f(sheet)
    ws_v = model.ws_v(sheet)
    if (ws_f.max_row or 0) < 1 or ws_f.max_column is None:
        return None

    # Prefer a defined Excel table (ListObject) when present.
    excel_table = None
    if ws_f.tables:
        name, ref = next(iter(ws_f.tables.items()))
        ref = ref if isinstance(ref, str) else ref.ref
        min_col, min_row, max_col, max_row = openpyxl.utils.cell.range_boundaries(ref)
        excel_table, header_row = name, min_row
        first_col, last_col = min_col, max_col
        last_row = max_row
        header_conf, issues = 1.0, []
    else:
        header_row, header_conf, issues = _guess_header_row(ws_f)
        if header_row is None:
            return None
        cols = [c for c in range(1, ws_f.max_column + 1) if not V.is_blank(ws_f.cell(header_row, c).value)]
        if not cols:
            return None
        first_col, last_col = cols[0], cols[-1]
        last_row = _last_data_row(ws_f, header_row, first_col, last_col)

    columns: list[Column] = []
    seen: dict[str, int] = {}
    for c in range(first_col, last_col + 1):
        raw = ws_f.cell(header_row, c).value
        name = str(raw).strip() if not V.is_blank(raw) else f"Column {get_column_letter(c)}"
        if V.is_blank(raw):
            issues.append(f"Column {get_column_letter(c)} has no header; named '{name}'.")
            header_conf = min(header_conf, 0.7)
        key = name.lower()
        if key in seen:
            seen[key] += 1
            issues.append(f"Duplicate header '{name}' in column {get_column_letter(c)}.")
            name = f"{name} ({seen[key]})"
            header_conf = min(header_conf, 0.7)
        else:
            seen[key] = 1
        columns.append(Column(index=c, letter=get_column_letter(c), name=name))

    rows: list[tuple[int, list[Any]]] = []
    formulas: dict[tuple[int, int], str] = {}
    for r in range(header_row + 1, last_row + 1):
        vals: list[Any] = []
        empty = True
        for col in columns:
            cell_f = ws_f.cell(r, col.index)
            raw = cell_f.value
            if cell_f.data_type == "f":
                formulas[(r, col.index)] = formula_text(raw)
                val = ws_v.cell(r, col.index).value
            else:
                val = raw
            if not V.is_blank(val) or (r, col.index) in formulas:
                empty = False
            vals.append(val)
        if not empty:
            rows.append((r, vals))

    table = SheetTable(
        sheet=sheet, header_row=header_row, first_data_row=header_row + 1, last_row=last_row,
        first_col=first_col, last_col=last_col, columns=columns, rows=rows, formulas=formulas,
        header_confidence=round(header_conf, 2), header_issues=issues, excel_table=excel_table,
    )
    from .profile import infer_column_types  # local import to avoid a cycle

    infer_column_types(table, ws_f)
    return table


def _guess_header_row(ws) -> tuple[int | None, float, list[str]]:
    """Pick the first of the top 10 rows that looks like a header (mostly text, followed by data)."""
    issues: list[str] = []
    max_col = min(ws.max_column or 0, 200)
    best = None
    for r in range(1, min(ws.max_row, 10) + 1):
        cells = [ws.cell(r, c).value for c in range(1, max_col + 1)]
        filled = [v for v in cells if not V.is_blank(v)]
        if len(filled) < 1:
            continue
        text_share = sum(isinstance(v, str) and not str(v).startswith("=") for v in filled) / len(filled)
        next_filled = [ws.cell(r + 1, c).value for c in range(1, max_col + 1)]
        has_data_below = any(not V.is_blank(v) for v in next_filled)
        if len(filled) >= 2 and text_share >= 0.8 and has_data_below:
            best = r
            break
        if best is None and has_data_below:
            best = r
    if best is None:
        return None, 0.0, ["No header row could be identified."]
    confidence = 1.0 if best == 1 else 0.85
    if best != 1:
        issues.append(f"Header row inferred at row {best} (rows above treated as titles/notes).")
    row_vals = [ws.cell(best, c).value for c in range(1, max_col + 1)]
    filled = [v for v in row_vals if not V.is_blank(v)]
    if filled and sum(isinstance(v, str) for v in filled) / len(filled) < 0.8:
        confidence = 0.5
        issues.append(f"Row {best} contains numbers or dates; it may be data rather than a header.")
    return best, confidence, issues


def _last_data_row(ws, header_row: int, first_col: int, last_col: int) -> int:
    last = header_row
    blank_run = 0
    for r in range(header_row + 1, (ws.max_row or header_row) + 1):
        if any(not V.is_blank(ws.cell(r, c).value) for c in range(first_col, last_col + 1)):
            last = r
            blank_run = 0
        else:
            blank_run += 1
            if blank_run >= 50:  # a long gap ends the table; later blocks are usually notes
                break
    return last


class ModelCache:
    """LRU cache of loaded models keyed by (workbook_id, version). Versions are immutable."""

    def __init__(self, capacity: int = 8):
        self.capacity = capacity
        self._items: OrderedDict[tuple[str, int], WorkbookModel] = OrderedDict()
        self._lock = threading.Lock()
        self.hits = 0
        self.misses = 0

    def get(self, key: tuple[str, int], loader) -> WorkbookModel:
        with self._lock:
            if key in self._items:
                self._items.move_to_end(key)
                self.hits += 1
                return self._items[key]
        model = loader()
        with self._lock:
            self.misses += 1
            self._items[key] = model
            while len(self._items) > self.capacity:
                self._items.popitem(last=False)
        return model

    def drop_workbook(self, workbook_id: str) -> None:
        with self._lock:
            for key in [k for k in self._items if k[0] == workbook_id]:
                del self._items[key]


model_cache = ModelCache()
