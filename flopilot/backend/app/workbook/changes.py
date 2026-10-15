"""Change sets: concrete, reviewable descriptions of a mutation, and the engine that applies them.

A change set lists every cell it will touch with its before/after state, plus any structural
edit (row deletion, new sheet). It is hashed; approval binds to that hash and to the workbook
version. Applying always starts from the immutable base version and produces new bytes;
nothing is modified in place.
"""
from __future__ import annotations

import copy
import io
from datetime import date, datetime
from typing import Any

import openpyxl
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.utils.cell import coordinate_from_string, column_index_from_string, range_boundaries
from openpyxl.worksheet.table import TableColumn

from ..common import AppError
from . import values as V
from .model import SheetTable, WorkbookModel

CHANGESET_SCHEMA = "1.0"
HEADER_FILL = "E2F0D9"


# --------------------------------------------------------------------------- value encoding
def encode(value: Any) -> dict:
    if value is None:
        return {"type": "blank", "value": None}
    if isinstance(value, bool):
        return {"type": "bool", "value": value}
    if V.is_number(value):
        return {"type": "number", "value": value}
    if isinstance(value, datetime):
        return {"type": "datetime", "value": value.isoformat()}
    if isinstance(value, date):
        return {"type": "date", "value": value.isoformat()}
    return {"type": "text", "value": str(value)}


def decode(enc: dict) -> Any:
    t, v = enc["type"], enc["value"]
    if t == "blank":
        return None
    if t == "date":
        return datetime.fromisoformat(v)
    if t == "datetime":
        return datetime.fromisoformat(v)
    return v


def cell_state(model: WorkbookModel, sheet: str, coord: str) -> dict:
    cell = model.ws_f(sheet)[coord]
    if cell.data_type == "f":
        return {"formula": str(getattr(cell.value, "text", cell.value)), **encode(model.ws_v(sheet)[coord].value)}
    return {"formula": None, **encode(cell.value)}


def value_change(model: WorkbookModel, sheet: str, coord: str, after: Any, number_format: str | None = None) -> dict:
    return {"sheet": sheet, "cell": coord, "before": cell_state(model, sheet, coord),
            "after": {"formula": None, **encode(after)}, "number_format": number_format}


def formula_change(model: WorkbookModel, sheet: str, coord: str, formula: str, number_format: str | None = None,
                   preview: Any = None) -> dict:
    return {"sheet": sheet, "cell": coord, "before": cell_state(model, sheet, coord),
            "after": {"formula": formula, **encode(preview)}, "number_format": number_format}


def summarize_flags(changes: list[dict]) -> dict:
    overwrites = sum(1 for c in changes if c["before"]["type"] != "blank" and c["before"]["formula"] is None
                     and (c["after"]["formula"] is not None or c["after"]["value"] != c["before"]["value"]))
    formula_replacements = sum(1 for c in changes if c["before"]["formula"] is not None)
    clears = sum(1 for c in changes if c["before"]["type"] != "blank" and c["after"]["type"] == "blank" and not c["after"]["formula"])
    return {"overwrites_values": overwrites, "replaces_formulas": formula_replacements, "clears_values": clears,
            "inserts_formulas": sum(1 for c in changes if c["after"]["formula"]), "deletes_rows": 0, "adds_sheet": False}


def make_change_set(action_type: str, title: str, description: str, sheet: str | None, changes: list[dict],
                    target_range: str | None = None, structural: dict | None = None, new_sheet: dict | None = None,
                    warnings: list[str] | None = None, extend_tables: dict | None = None,
                    details: dict | None = None) -> dict:
    flags = summarize_flags(changes)
    if structural and structural.get("delete_rows"):
        flags["deletes_rows"] = len(structural["delete_rows"])
    if new_sheet:
        flags["adds_sheet"] = True
    affected = len(changes)
    if structural and structural.get("delete_rows"):
        affected += len(structural["delete_rows"]) * structural.get("row_width", 1)
    if new_sheet:
        affected += sum(len(r) for r in new_sheet["rows"]) + len(new_sheet["header"])
    destructive = bool(flags["deletes_rows"] or flags["replaces_formulas"] or flags["clears_values"])
    return {
        "schema_version": CHANGESET_SCHEMA,
        "action_type": action_type,
        "title": title,
        "description": description,
        "sheet": sheet,
        "target_range": target_range,
        "changes": changes,
        "structural": structural,
        "new_sheet": new_sheet,
        "extend_tables": extend_tables,
        "flags": flags,
        "destructive": destructive,
        "affected_cells": affected,
        "warnings": warnings or [],
        "details": details or {},
        "preconditions": [
            "Workbook version is unchanged since the proposal.",
            "Each target cell still holds its 'before' value.",
        ],
    }


# --------------------------------------------------------------------------- builders
def build_formula_column(model: WorkbookModel, proposal: dict, previews: dict[int, Any]) -> dict:
    sheet = proposal["sheet"]
    letter = proposal["target_column"]
    table = model.table(sheet)
    changes = [value_change(model, sheet, proposal["header_cell"], proposal["new_column"])]
    for r in proposal["rows"]:
        coord = f"{letter}{r}"
        changes.append(formula_change(model, sheet, coord, proposal["template"].replace("{r}", str(r)),
                                      proposal["number_format"], previews.get(r)))
    extend = None
    if table is not None and column_index_from_string(letter) == table.last_col + 1:
        extend = {"sheet": sheet, "column": letter, "header": proposal["new_column"], "header_row": table.header_row}
    return make_change_set(
        "add_formula_column", f"Add '{proposal['new_column']}' column",
        f"Insert a header in {proposal['header_cell']} and the formula {proposal['first_formula']} (adjusted per row) in {proposal['target_range']}.",
        sheet, changes, proposal["target_range"], extend_tables=extend,
        details={"formula_proposal": proposal},
    )


def build_set_formula(model: WorkbookModel, sheet: str, coord: str, formula: str, reason: str) -> dict:
    change = formula_change(model, sheet, coord, formula, None)
    return make_change_set("set_formula", f"Replace formula in {sheet}!{coord}", reason, sheet, [change], coord,
                           warnings=["This replaces an existing formula."] if change["before"]["formula"] else [])


def build_summary_sheet(model: WorkbookModel, result: dict, title: str, chart: dict | None) -> dict:
    """A new sheet holding a computed summary table (static values) and optionally a native chart."""
    base = (title or "Summary")[:24].replace("/", "-").replace("\\", "-").replace("?", "").replace("*", "").replace("[", "(").replace("]", ")").replace(":", "-")
    name, n = base, 2
    existing = {s.lower() for s in model.sheet_names}
    while name.lower() in existing:
        name, n = f"{base} {n}"[:31], n + 1
    header, rows = summary_table(result)
    new_sheet = {
        "name": name, "title": title, "header": header, "rows": rows,
        "note": f"Computed by FloPilot from {result.get('sheet')} (workbook version {model.version}). Static values: re-run to refresh.",
        "number_format": result.get("number_format") or "#,##0.00",
        "chart": chart,
    }
    return make_change_set("create_summary_sheet", f"Create report sheet '{name}'",
                           f"Add a new sheet '{name}' with {len(rows)} summary row(s)" + (" and a native Excel chart." if chart else "."),
                           name, [], f"A1:{get_column_letter(len(header))}{len(rows) + 4}", new_sheet=new_sheet)


def summary_table(result: dict) -> tuple[list[str], list[list[Any]]]:
    kind = result.get("kind")
    if kind == "time_series":
        header = ["Period"] + list(result["series"].keys())
        rows = [[p] + [next((pt["value"] for pt in result["series"][g] if pt["period"] == p), None) for g in result["series"]]
                for p in result["periods"]]
        return header, rows
    if kind == "variance":
        header = (result["group_by"] or ["Group"]) + [result["actual"], result["budget"], "Variance", "Variance %"]
        return header, [r["group"] + [r["actual"], r["budget"], r["value"], (r["pct"] / 100 if r["pct"] is not None else None)] for r in result["rows"]]
    if kind == "period_change":
        header = result["group_by"] + [result["previous_period"], result["current_period"], "Change", "Change %"]
        return header, [r["group"] + [r["previous"], r["current"], r["value"], (r["pct"] / 100 if r["pct"] is not None else None)] for r in result["rows"]]
    label = f"{(result.get('function') or 'value').title()} of {result.get('measure') or 'rows'}"
    header = (result.get("group_by") or ["Group"]) + [label]
    return header, [r["group"] + [r["value"]] for r in result["rows"]]


def build_flag_column(model: WorkbookModel, table: SheetTable, header: str, flags: dict[int, str], reason: str) -> dict:
    from .formulas import next_free_column

    col = next_free_column(table, model)
    letter = get_column_letter(col)
    changes = [value_change(model, table.sheet, f"{letter}{table.header_row}", header)]
    for r, label in sorted(flags.items()):
        changes.append(value_change(model, table.sheet, f"{letter}{r}", label))
    extend = {"sheet": table.sheet, "column": letter, "header": header, "header_row": table.header_row} if col == table.last_col + 1 else None
    return make_change_set("add_flag_column", f"Add '{header}' review column", reason, table.sheet, changes,
                           f"{letter}{table.header_row}:{letter}{table.last_row}", extend_tables=extend)


def build_delete_rows(model: WorkbookModel, table: SheetTable, rows: list[int], reason: str) -> dict:
    rows = sorted(set(rows))
    problems = structural_edit_blockers(model, table.sheet, min(rows))
    if problems:
        raise AppError("structural_edit_blocked", "Rows cannot be safely deleted: " + " ".join(problems)
                       + " Use 'flag duplicates' instead.", 409, {"blockers": problems})
    width = table.last_col - table.first_col + 1
    removed = []
    for r in rows:
        vals = next(v for rr, v in table.rows if rr == r)
        removed.append({"row": r, "values": [encode(v) for v in vals]})
    return make_change_set("delete_rows", f"Delete {len(rows)} row(s) from {table.sheet}", reason, table.sheet, [],
                           f"{table.sheet} rows {', '.join(map(str, rows[:10]))}{'…' if len(rows) > 10 else ''}",
                           structural={"delete_rows": rows, "row_width": width, "removed": removed,
                                       "header_row": table.header_row},
                           warnings=["Rows will be permanently removed from the new version (earlier versions remain restorable)."])


def structural_edit_blockers(model: WorkbookModel, sheet: str, first_row: int) -> list[str]:
    from .evaluator import parse_ref, referenced_ranges

    problems = []
    ws = model.ws_f(sheet)
    if any(rng.min_row >= first_row or rng.max_row >= first_row for rng in ws.merged_cells.ranges):
        problems.append("The sheet has merged cells below the first deleted row.")
    refs = 0
    for other in model.wb_f.worksheets:
        for row in other.iter_rows():
            for cell in row:
                if cell.data_type != "f":
                    continue
                for ref in referenced_ranges(str(getattr(cell.value, "text", cell.value))):
                    try:
                        s, r1, _, r2, _ = parse_ref(ref, other.title)
                    except Exception:
                        continue
                    if s == sheet and r2 >= first_row and not (other.title == sheet and cell.row >= first_row and r1 == r2 == cell.row):
                        refs += 1
    if refs:
        problems.append(f"{refs} formula reference(s) point at rows that would shift.")
    if ws.data_validations and ws.data_validations.dataValidation:
        problems.append("The sheet has data-validation rules that openpyxl would not shift.")
    if list(ws.conditional_formatting):
        problems.append("The sheet has conditional formatting that openpyxl would not shift.")
    return problems


# --------------------------------------------------------------------------- apply engine
def _write(cell, after: dict, number_format: str | None) -> None:
    if after.get("formula"):
        cell.value = after["formula"]
    else:
        value = decode(after)
        cell.value = value
        if isinstance(value, str):
            cell.data_type = "s"  # text is never interpreted as a formula
        if isinstance(value, datetime) and not number_format:
            cell.number_format = "yyyy-mm-dd"
    if number_format:
        cell.number_format = number_format


def _matches_before(ws_cell, ws_cached_value, before: dict) -> bool:
    if before["formula"] is not None:
        return ws_cell.data_type == "f" and str(getattr(ws_cell.value, "text", ws_cell.value)) == before["formula"]
    current = encode(ws_cell.value)
    if current["type"] in ("date", "datetime") and before["type"] in ("date", "datetime"):
        return str(current["value"])[:10] == str(before["value"])[:10]
    return current["type"] == before["type"] and current["value"] == before["value"]


def apply_change_set(base: bytes, cs: dict) -> bytes:
    if cs.get("schema_version") != CHANGESET_SCHEMA:
        raise AppError("unsupported_change_set", "This change set uses an unsupported schema version.")
    wb = openpyxl.load_workbook(io.BytesIO(base), keep_links=True)
    wb_v = openpyxl.load_workbook(io.BytesIO(base), data_only=True, keep_links=True)
    styled_cols: set[tuple[str, str]] = set()
    for change in cs["changes"]:
        if change["sheet"] not in wb.sheetnames:
            raise AppError("precondition_failed", f"Sheet '{change['sheet']}' no longer exists.", 409)
        ws = wb[change["sheet"]]
        cell = ws[change["cell"]]
        if not _matches_before(cell, wb_v[change["sheet"]][change["cell"]].value, change["before"]):
            raise AppError("precondition_failed", f"{change['sheet']}!{change['cell']} changed since the proposal was made.", 409)
        col_letter, row = coordinate_from_string(change["cell"])
        # Only cells that were empty and unstyled (i.e. new cells) take their neighbour's look.
        if not cell.has_style and change["before"]["type"] == "blank" and change["before"]["formula"] is None:
            _copy_neighbor_style(ws, cell, col_letter, row)
            styled_cols.add((change["sheet"], col_letter))
        _write(cell, change["after"], change.get("number_format"))

    for sheet, letter in styled_cols:
        ws = wb[sheet]
        dim = ws.column_dimensions.get(letter)
        if dim is None or not dim.customWidth:
            prev = get_column_letter(max(column_index_from_string(letter) - 1, 1))
            ws.column_dimensions[letter].width = max(ws.column_dimensions[prev].width or 12, 14)

    if cs.get("extend_tables"):
        _extend_tables(wb, cs["extend_tables"])
    if cs.get("structural") and cs["structural"].get("delete_rows"):
        _delete_rows(wb[cs["sheet"]], cs["structural"])
    if cs.get("new_sheet"):
        _create_sheet(wb, cs["new_sheet"])
    # openpyxl drops cached results, so ask Excel to recalculate everything when the file is opened.
    wb.calculation.fullCalcOnLoad = True
    out = io.BytesIO()
    wb.save(out)
    from .cache_values import refresh_cached_values

    candidate, stats = refresh_cached_values(base, out.getvalue(), cs)
    cs.setdefault("details", {})["cached_values"] = stats
    return candidate


def _copy_neighbor_style(ws, cell, col_letter: str, row: int) -> None:
    """Give a newly written cell the look of its left neighbour (font, fill, border, alignment)."""
    idx = column_index_from_string(col_letter)
    if idx <= 1:
        return
    left = ws.cell(row=row, column=idx - 1)
    if left.has_style:
        cell.font = copy.copy(left.font)
        cell.fill = copy.copy(left.fill)
        cell.border = copy.copy(left.border)
        cell.alignment = copy.copy(left.alignment)


def _extend_tables(wb, ext: dict) -> None:
    ws = wb[ext["sheet"]]
    new_col = column_index_from_string(ext["column"])
    for table in ws.tables.values():
        min_col, min_row, max_col, max_row = range_boundaries(table.ref)
        if max_col + 1 == new_col and min_row == ext["header_row"]:
            table.ref = f"{get_column_letter(min_col)}{min_row}:{ext['column']}{max_row}"
            table.tableColumns.append(TableColumn(id=len(table.tableColumns) + 1, name=ext["header"]))
            if table.autoFilter is not None:
                table.autoFilter.ref = table.ref
    if ws.auto_filter and ws.auto_filter.ref:
        min_col, min_row, max_col, max_row = range_boundaries(ws.auto_filter.ref)
        if max_col + 1 == new_col:
            ws.auto_filter.ref = f"{get_column_letter(min_col)}{min_row}:{ext['column']}{max_row}"


def _delete_rows(ws, structural: dict) -> None:
    rows = sorted(structural["delete_rows"], reverse=True)
    heights = {r: d.height for r, d in ws.row_dimensions.items() if d.height}
    for r in rows:
        ws.delete_rows(r, 1)
    # shift custom row heights to follow their rows
    for r in sorted(heights):
        shift = sum(1 for d in rows if d < r)
        if shift and r not in rows:
            ws.row_dimensions[r - shift].height = heights[r]
    n = len(rows)
    for table in ws.tables.values():
        min_col, min_row, max_col, max_row = range_boundaries(table.ref)
        table.ref = f"{get_column_letter(min_col)}{min_row}:{get_column_letter(max_col)}{max_row - n}"
        if table.autoFilter is not None:
            table.autoFilter.ref = table.ref
    if ws.auto_filter and ws.auto_filter.ref:
        min_col, min_row, max_col, max_row = range_boundaries(ws.auto_filter.ref)
        ws.auto_filter.ref = f"{get_column_letter(min_col)}{min_row}:{get_column_letter(max_col)}{max_row - n}"


def _create_sheet(wb, spec: dict) -> None:
    ws = wb.create_sheet(spec["name"])
    title = ws.cell(row=1, column=1, value=spec.get("title") or spec["name"])
    title.data_type = "s"
    title.font = Font(bold=True, size=14, color="0B4F35")
    note = ws.cell(row=2, column=1, value=spec["note"])
    note.data_type = "s"
    note.font = Font(italic=True, color="6B7280", size=9)
    header_row = 4
    for c, h in enumerate(spec["header"], start=1):
        cell = ws.cell(row=header_row, column=c, value=str(h))
        cell.data_type = "s"
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor=HEADER_FILL)
        cell.alignment = Alignment(horizontal="center")
        ws.column_dimensions[get_column_letter(c)].width = max(14, min(40, len(str(h)) + 4))
    for r, row in enumerate(spec["rows"], start=header_row + 1):
        for c, value in enumerate(row, start=1):
            cell = ws.cell(row=r, column=c, value=value)
            if isinstance(value, str):
                cell.data_type = "s"
            elif V.is_number(value):
                header = str(spec["header"][c - 1])
                cell.number_format = "0.0%" if header.endswith("%") else spec["number_format"]
    ws.freeze_panes = f"A{header_row + 1}"
    chart_spec = spec.get("chart")
    if chart_spec and spec["rows"]:
        n_rows = len(spec["rows"])
        chart = LineChart() if chart_spec.get("type") == "line" else BarChart()
        chart.title = chart_spec.get("title") or spec["name"]
        chart.y_axis.title = chart_spec.get("y_title") or ""
        chart.x_axis.title = chart_spec.get("x_title") or ""
        chart.height, chart.width = 8, 16
        last_col = 2 if chart_spec.get("type") != "line" else len(spec["header"])
        data = Reference(ws, min_col=2, max_col=last_col, min_row=header_row, max_row=header_row + n_rows)
        cats = Reference(ws, min_col=1, min_row=header_row + 1, max_row=header_row + n_rows)
        chart.add_data(data, titles_from_data=True)
        chart.set_categories(cats)
        ws.add_chart(chart, f"{get_column_letter(len(spec['header']) + 2)}{header_row}")
