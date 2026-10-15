"""Workbook profiling: structure, inferred types, missing values and data-quality issues.

Profiles are computed deterministically and cached per immutable version, so the LLM (when
configured) only ever sees this compact schema rather than the whole workbook.
"""
from __future__ import annotations

import re
import statistics
from collections import Counter
from typing import Any

from openpyxl.utils import get_column_letter

from . import values as V
from .model import Column, SheetTable, WorkbookModel
from .quality import category_variants, date_format_mix, find_duplicates, outlier_candidates, whitespace_cells

PROFILE_SAMPLE_ROWS = 5000

_CURRENCY_WORDS = re.compile(r"revenue|cost|price|amount|sales|spend|budget|actual|profit|value|salary|total|fee|expense", re.I)
_PERCENT_WORDS = re.compile(r"margin|rate|percent|pct|%|discount|share", re.I)
_ID_WORDS = re.compile(r"(^id$|\bid\b|_id$| id$|number$|\bno\.?$|code$|#$|^sku$|invoice)", re.I)


def infer_column_types(table: SheetTable, ws) -> None:
    sample = table.rows[:PROFILE_SAMPLE_ROWS]
    for col in table.columns:
        pos = table.col_pos(col)
        kinds = Counter(V.classify(vals[pos]) for _, vals in sample)
        kinds.pop("blank", None)
        total = sum(kinds.values())
        fmt = None
        for r, _ in sample[:20]:
            cell_fmt = ws.cell(r, col.index).number_format
            if cell_fmt and cell_fmt != "General":
                fmt = cell_fmt
                break
        col.number_format = fmt
        if total == 0:
            col.inferred_type, col.type_confidence = "empty", 1.0
            continue
        merged = Counter()
        for kind, n in kinds.items():
            merged[{"date_text": "date", "numeric_text": "number"}.get(kind, kind)] += n
        merged.pop("error", None)
        dominant = merged.most_common(1)[0][0] if merged else "error"
        raw_kinds = Counter({k: n for k, n in kinds.items() if k != "error"})
        col.inferred_type = dominant
        col.type_confidence = round(raw_kinds.most_common(1)[0][1] / total, 2) if raw_kinds else 0.0
        if len([k for k in kinds if k != "error"]) > 1:
            col.inferred_type = f"mixed ({dominant})"
        col.semantic = _semantic(col, table, dominant, fmt)


def _semantic(col: Column, table: SheetTable, dominant: str, fmt: str | None) -> str | None:
    name = col.name
    vals = [v for v in table.values(col) if not V.is_blank(v)]
    distinct = len({str(v).strip().lower() for v in vals})
    if dominant == "date":
        return "date"
    if dominant == "number":
        if fmt and "%" in fmt:
            return "percentage"
        if fmt and any(sym in fmt for sym in "$€£¥"):
            return "currency"
        if _ID_WORDS.search(name) and distinct >= 0.9 * len(vals):
            return "identifier"
        if _PERCENT_WORDS.search(name) and all(-1.5 <= float(v) <= 1.5 for v in vals if V.is_number(v)):
            return "percentage"
        if _CURRENCY_WORDS.search(name):
            return "currency"
        return "numeric"
    if dominant == "text":
        if _ID_WORDS.search(name) and vals and distinct >= 0.9 * len(vals):
            return "identifier"
        if vals and distinct <= max(50, 0.5 * len(vals)):
            return "category"
        return "text"
    return None


def profile_sheet(model: WorkbookModel, sheet: str, deep: bool = True) -> dict:
    ws = model.ws_f(sheet)
    ws_v = model.ws_v(sheet)
    table = model.table(sheet)
    hidden_rows = [r for r, dim in ws.row_dimensions.items() if dim.hidden]
    hidden_cols = [k for k, dim in ws.column_dimensions.items() if dim.hidden]

    formula_cells = 0
    uncached = 0
    error_cells: list[dict] = []
    for row in ws.iter_rows():
        for cell in row:
            if cell.data_type == "f":
                formula_cells += 1
                cached = ws_v[cell.coordinate].value
                if cached is None:
                    uncached += 1
                elif V.is_error(cached) and len(error_cells) < 25:
                    error_cells.append({"cell": cell.coordinate, "error": cached, "formula": str(getattr(cell.value, "text", cell.value))})
            elif V.is_error(cell.value) and len(error_cells) < 25:
                error_cells.append({"cell": cell.coordinate, "error": cell.value, "formula": None})

    info: dict[str, Any] = {
        "name": sheet,
        "state": ws.sheet_state,
        "dimensions": ws.dimensions,
        "max_row": ws.max_row,
        "max_column": ws.max_column,
        "merged_ranges": [str(r) for r in ws.merged_cells.ranges][:50],
        "hidden_rows": len(hidden_rows),
        "hidden_columns": len(hidden_cols),
        "freeze_panes": ws.freeze_panes,
        "auto_filter": ws.auto_filter.ref if ws.auto_filter and ws.auto_filter.ref else None,
        "excel_tables": list(ws.tables.keys()),
        "data_validations": len(ws.data_validations.dataValidation) if ws.data_validations else 0,
        "conditional_formats": len(list(ws.conditional_formatting)),
        "formula_cells": formula_cells,
        "formulas_without_cached_value": uncached,
        "error_cells": error_cells,
        "table": None,
        "issues": [],
    }
    if table is None:
        info["issues"].append(_issue("empty_sheet", "info", sheet, None, "This sheet has no tabular data.", 0))
        return info

    columns = []
    for col in table.columns:
        vals = table.values(col)
        non_blank = [v for v in vals if not V.is_blank(v)]
        missing = len(vals) - len(non_blank)
        entry: dict[str, Any] = {
            "name": col.name, "letter": col.letter, "inferred_type": col.inferred_type, "semantic": col.semantic,
            "type_confidence": col.type_confidence, "number_format": col.number_format,
            "missing": missing, "missing_pct": round(100 * missing / len(vals), 1) if vals else 0.0,
            "distinct": len({str(v) for v in non_blank}),
            "sample": [V.json_safe(v) for v in non_blank[:3]],
            "range": table.range_ref(col),
        }
        nums = [float(v) for v in non_blank if V.is_number(v)]
        if nums and "number" in col.inferred_type:
            entry["stats"] = {"min": min(nums), "max": max(nums), "mean": round(statistics.fmean(nums), 4), "sum": round(sum(nums), 4), "count": len(nums)}
        columns.append(entry)

    info["table"] = {
        "header_row": table.header_row, "first_data_row": table.first_data_row, "last_row": table.last_row,
        "range": table.range_ref(), "row_count": len(table.rows), "column_count": len(table.columns),
        "header_confidence": table.header_confidence, "header_issues": table.header_issues,
        "excel_table": table.excel_table, "columns": columns,
    }
    if deep:
        info["issues"] = _issues(model, table, info)
    return info


def _issue(kind: str, severity: str, sheet: str, column: str | None, message: str, count: int,
           range_ref: str | None = None, suggestion: str | None = None, **extra: Any) -> dict:
    return {"id": f"{sheet}:{column or '*'}:{kind}", "kind": kind, "severity": severity, "sheet": sheet,
            "column": column, "message": message, "count": count, "range": range_ref, "suggested_prompt": suggestion, **extra}


def _issues(model: WorkbookModel, table: SheetTable, info: dict) -> list[dict]:
    sheet = table.sheet
    issues: list[dict] = []
    if table.header_confidence < 0.9:
        issues.append(_issue("uncertain_header", "warning", sheet, None,
                             "Header detection is uncertain: " + "; ".join(table.header_issues), 1))
    for entry, col in zip(info["table"]["columns"], table.columns):
        if entry["missing"]:
            issues.append(_issue("missing_values", "warning" if entry["missing_pct"] >= 5 else "info", sheet, col.name,
                                 f"{entry['missing']} missing value(s) in '{col.name}' ({entry['missing_pct']}%).",
                                 entry["missing"], entry["range"], f"Show missing values by column on {sheet}"))
        if col.type_confidence < 0.95 or col.inferred_type.startswith("mixed"):
            issues.append(_issue("mixed_types", "warning", sheet, col.name,
                                 f"'{col.name}' mixes data types (inferred {col.inferred_type}, {int(col.type_confidence * 100)}% consistent).",
                                 1, entry["range"]))
        if col.semantic == "date" or col.inferred_type.endswith("(date)"):
            mix = date_format_mix(table, col)
            text_kinds = {k: n for k, n in mix["kinds"].items() if k.startswith("text:")}
            if text_kinds:
                n = sum(text_kinds.values())
                issues.append(_issue("inconsistent_dates", "warning", sheet, col.name,
                                     f"{n} date(s) in '{col.name}' are stored as text in {len(text_kinds)} format(s).",
                                     n, entry["range"], f"Standardize the dates in {col.name} on {sheet}", detail=mix))
        if col.semantic == "category":
            groups = category_variants(table, col)
            if groups:
                n = sum(sum(g["variants"].values()) for g in groups)
                issues.append(_issue("category_variants", "warning", sheet, col.name,
                                     f"{n} value(s) in '{col.name}' look like spelling variants of {len(groups)} categor{'y' if len(groups) == 1 else 'ies'}.",
                                     n, entry["range"], f"Normalize the {col.name} categories on {sheet}",
                                     detail=[{"canonical": g["canonical"], "variants": list(g["variants"])} for g in groups[:10]]))
            ws_rows = whitespace_cells(table, col)
            if ws_rows:
                issues.append(_issue("whitespace", "info", sheet, col.name,
                                     f"{len(ws_rows)} value(s) in '{col.name}' have extra spaces.", len(ws_rows), entry["range"],
                                     f"Trim whitespace in {col.name} on {sheet}"))
        if col.semantic in ("currency", "numeric") and len(table.rows) >= 8:
            out = outlier_candidates(table, col)
            if out["candidates"]:
                issues.append(_issue("outlier_candidates", "info", sheet, col.name,
                                     f"{len(out['candidates'])} unusually large or small value(s) in '{col.name}' (candidates for review).",
                                     len(out["candidates"]), entry["range"], f"Identify unusually large transactions in {sheet}",
                                     rows=[c["row"] for c in out["candidates"][:10]]))
    dups = find_duplicates(table)
    if dups["groups"]:
        issues.append(_issue("duplicate_candidates", "warning", sheet, None,
                             f"{dups['extra_rows']} row(s) look like duplicates of other rows (matching on {len(dups['compared_columns'])} columns).",
                             dups["extra_rows"], table.range_ref(), f"Find potential duplicates on {sheet}",
                             rows=[g for g in dups["groups"][:10]]))
    for rep in dups["repeated_identifiers"][:5]:
        issues.append(_issue("repeated_identifier", "warning", sheet, rep["column"],
                             f"Identifier '{rep['value']}' in '{rep['column']}' appears {len(rep['rows'])} times (rows {', '.join(map(str, rep['rows']))}).",
                             len(rep["rows"]) - 1, None, f"Find potential duplicates on {sheet}"))
    for err in info["error_cells"]:
        issues.append(_issue("formula_error", "error", sheet, None,
                             f"{err['cell']} shows {err['error']}" + (f" from {err['formula']}" if err["formula"] else "") + ".",
                             1, err["cell"], f"Diagnose the formula error in {sheet}!{err['cell']}", cell=err["cell"]))
    if info["formulas_without_cached_value"]:
        issues.append(_issue("uncalculated_formulas", "info", sheet, None,
                             f"{info['formulas_without_cached_value']} formula(s) have no saved result; values appear blank until recalculated in Excel.",
                             info["formulas_without_cached_value"]))
    if info["merged_ranges"]:
        issues.append(_issue("merged_cells", "info", sheet, None,
                             f"{len(info['merged_ranges'])} merged range(s); structural edits on this sheet are restricted.", len(info["merged_ranges"])))
    if info["hidden_rows"] or info["hidden_columns"]:
        issues.append(_issue("hidden_content", "info", sheet, None,
                             f"{info['hidden_rows']} hidden row(s) and {info['hidden_columns']} hidden column(s). Calculations include hidden rows.",
                             info["hidden_rows"] + info["hidden_columns"]))
    return issues


def profile_workbook(model: WorkbookModel) -> dict:
    sheets = [profile_sheet(model, s) for s in model.sheet_names]
    issues = [i for s in sheets for i in s["issues"]]
    hidden = [s["name"] for s in sheets if s["state"] != "visible"]
    if hidden:
        issues.append(_issue("hidden_sheets", "info", hidden[0], None, f"Hidden sheet(s): {', '.join(hidden)}.", len(hidden)))
    from .security import find_instruction_like_text

    for hit in find_instruction_like_text(model):
        issues.append(_issue("instruction_like_text", "info", hit["sheet"], None,
                             f"{hit['sheet']}!{hit['cell']} contains text that reads like instructions to an AI assistant. "
                             "FloPilot treats workbook content as data and will not follow it.", 1, hit["cell"]))
    severity_order = {"error": 0, "warning": 1, "info": 2}
    issues.sort(key=lambda i: (severity_order[i["severity"]], i["sheet"]))
    return {
        "schema_version": "1.0",
        "workbook_id": model.workbook_id,
        "version": model.version,
        "sheet_count": len(sheets),
        "sheets": sheets,
        "issues": issues,
        "issue_counts": dict(Counter(i["severity"] for i in issues)),
        "defined_names": sorted(model.wb_f.defined_names.keys()) if hasattr(model.wb_f.defined_names, "keys") else [],
    }


def compact_schema(profile: dict, max_columns: int = 40) -> dict:
    """Bounded schema summary for the LLM: no cell values beyond 3 samples per column."""
    out = []
    for s in profile["sheets"]:
        t = s.get("table")
        out.append({
            "sheet": s["name"], "state": s["state"],
            "rows": t["row_count"] if t else 0,
            "range": t["range"] if t else None,
            "columns": [
                {k: c[k] for k in ("name", "letter", "inferred_type", "semantic", "missing")}
                for c in (t["columns"][:max_columns] if t else [])
            ],
        })
    return {"sheets": out, "issue_count": len(profile["issues"])}


def column_letter(index: int) -> str:
    return get_column_letter(index)
