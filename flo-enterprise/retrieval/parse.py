"""Parsers producing (section, text) blocks for PDF, DOCX, XLSX, CSV and Markdown, plus structured financial records."""
from __future__ import annotations

import csv
import io
import re

import yaml


def parse_markdown(text: str) -> tuple[dict, list[tuple[str, str]]]:
    meta: dict = {}
    if text.startswith("---"):
        _, fm, text = text.split("---", 2)
        meta = yaml.safe_load(fm) or {}
    sections, cur, buf = [], "Introduction", []
    for line in text.splitlines():
        m = re.match(r"^#{1,6}\s+(.*)", line)
        if m:
            if "".join(buf).strip():
                sections.append((cur, "\n".join(buf).strip()))
            cur, buf = m.group(1).strip(), []
        else:
            buf.append(line)
    if "".join(buf).strip():
        sections.append((cur, "\n".join(buf).strip()))
    return meta, sections


def parse_bytes(filename: str, data: bytes) -> tuple[dict, list[tuple[str, str]]]:
    ext = filename.rsplit(".", 1)[-1].lower()
    if ext in ("md", "txt"):
        return parse_markdown(data.decode("utf-8"))
    if ext == "csv":
        rows = list(csv.reader(io.StringIO(data.decode("utf-8"))))
        header, body = rows[0], rows[1:]
        return {}, [(f"Row {i + 1}", "; ".join(f"{h}: {v}" for h, v in zip(header, r))) for i, r in enumerate(body)]
    if ext == "xlsx":
        from openpyxl import load_workbook
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        out = []
        for ws in wb.worksheets:
            rows = list(ws.iter_rows(values_only=True))
            if not rows:
                continue
            header = [str(h) for h in rows[0]]
            for i, r in enumerate(rows[1:]):
                out.append((f"{ws.title} row {i + 1}", "; ".join(f"{h}: {v}" for h, v in zip(header, r) if v is not None)))
        return {}, out
    if ext == "docx":
        from docx import Document
        doc = Document(io.BytesIO(data))
        sections, cur, buf = [], "Introduction", []
        for p in doc.paragraphs:
            if p.style.name.lower().startswith("heading"):
                if buf:
                    sections.append((cur, "\n".join(buf)))
                cur, buf = p.text.strip(), []
            elif p.text.strip():
                buf.append(p.text.strip())
        if buf:
            sections.append((cur, "\n".join(buf)))
        return {}, sections
    if ext == "pdf":
        from pypdf import PdfReader
        r = PdfReader(io.BytesIO(data))
        return {}, [(f"Page {i + 1}", (pg.extract_text() or "").strip()) for i, pg in enumerate(r.pages) if (pg.extract_text() or "").strip()]
    raise ValueError(f"unsupported type {ext}")


def records_to_sections(name: str, rows: list[dict], key: str) -> list[tuple[str, str]]:
    """Structured financial records (e.g. fund facts) indexed as one section per record."""
    return [(f"{name} {r[key]}", "; ".join(f"{k}: {v}" for k, v in r.items())) for r in rows]


def chunk(sections: list[tuple[str, str]], size: int = 700, overlap: int = 100) -> list[tuple[str, str]]:
    out = []
    for sec, text in sections:
        if len(text) <= size:
            out.append((sec, text))
            continue
        start = 0
        while start < len(text):
            end = min(len(text), start + size)
            cut = text.rfind(". ", start + size // 2, end)
            end = cut + 1 if cut != -1 and end < len(text) else end
            out.append((sec, text[start:end].strip()))
            if end >= len(text):
                break
            start = max(end - overlap, start + 1)
    return out
