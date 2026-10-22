import io
from datetime import date
from decimal import Decimal

import pytest

from backend.db.session import session_scope
from data.synthetic.generator import SYNTHETIC_LABEL, TARGET_COUNTS, generate, manifest
from integrations.mock.adapters import Adapters


@pytest.fixture(scope="module")
def ds():
    return generate()


@pytest.mark.critical
def test_generator_is_deterministic(ds):
    assert manifest(ds) == manifest(generate())
    assert "SYNTHETIC" in SYNTHETIC_LABEL


@pytest.mark.critical
def test_required_counts(ds):
    for k, v in TARGET_COUNTS.items():
        if k == "months":
            assert len({r["month"] for r in ds["product_monthly"]}) == v
        else:
            assert len(ds[k]) == v, k


@pytest.mark.critical
def test_referential_integrity(ds):
    ids = {k: {r["id"] for r in ds[k]} for k in ("clients", "advisors", "portfolios", "products", "securities", "benchmarks", "controls")}
    assert all(p["client_id"] in ids["clients"] for p in ds["portfolios"])
    assert all(c["advisor_id"] in ids["advisors"] for c in ds["clients"])
    assert all(t["portfolio_id"] in ids["portfolios"] and t["client_id"] in ids["clients"] for t in ds["transactions"])
    assert all(h["security_id"] in ids["securities"] for h in ds["portfolio_holdings"] + ds["product_holdings"])
    assert all(e["control_id"] in ids["controls"] for e in ds["compliance_events"])
    cl = {c["id"]: c for c in ds["clients"]}
    assert all(cl[t["client_id"]]["tenant_id"] == t["tenant_id"] for t in ds["transactions"])


def test_pii_is_obviously_synthetic(ds):
    assert all(c["email"].endswith("@example.com") for c in ds["clients"])
    assert all(c["sin"].startswith("0") for c in ds["clients"])
    assert all("-555-01" in c["phone"] for c in ds["clients"])


def test_narrative_aum_up_margin_down(st):
    from analytics.service import executive_kpis
    k = executive_kpis(Adapters(st, "t-flo"), 12)["data"]
    assert k["aum_growth"] > 0 and k["operating_margin"] < k["prior_period"]["operating_margin"]
    assert k["revenue_reconciliation"]["balanced"]


def test_retrieval_acl_tenant_expiry_quarantine(st, engine, P):
    from retrieval.search import search
    with session_scope() as s:
        ids = lambda u, q: {r["document_id"] for r in search(s, engine, P[u], q, 10)["results"]}  # noqa: E731
        assert "DOC-010" in ids("u-ceo", "ETF shelf expansion strategy")
        assert "DOC-010" not in ids("u-advisor", "ETF shelf expansion strategy")       # executive-only group
        assert "DOC-006" not in ids("u-advisor", "AML large cash transactions")         # restricted
        assert "DOC-006" in ids("u-compliance", "AML large cash transactions")
        assert "DOC-016" not in ids("u-advisor", "Northwind concentration limit")      # other tenant
        assert "DOC-016" in ids("u-nw-advisor", "Northwind concentration limit")
        assert "DOC-015" not in ids("u-cio", "trading desk guideline concentration")   # expired
        assert "DOC-014" not in ids("u-cio", "emerging markets appendix instructions")  # quarantined
        assert ids("u-external", "concentration policy suitability") <= {"DOC-009", "DOC-022"}  # public only


def test_multilingual_retrieval(st, engine, P):
    from retrieval.search import search
    with session_scope() as s:
        fr = search(s, engine, P["u-advisor"], "approbation des communications par la conformité", 3)
        es = search(s, engine, P["u-cio"], "¿Qué riesgo de liquidez tienen los mercados emergentes?", 3)
    assert fr["language"] == "fr" and fr["results"][0]["document_id"] == "DOC-013"
    assert es["language"] == "es" and "DOC-019" in {r["document_id"] for r in es["results"]}


def _docx():
    from docx import Document
    d = Document()
    d.add_heading("Fee schedule", 1)
    d.add_paragraph("Management fee is 0.35% for the institutional pool.")
    b = io.BytesIO()
    d.save(b)
    return b.getvalue()


def _xlsx():
    from openpyxl import Workbook
    wb = Workbook()
    ws = wb.active
    ws.append(["fund", "fee_bps"])
    ws.append(["FLO Pool", 35])
    b = io.BytesIO()
    wb.save(b)
    return b.getvalue()


def _pdf(text="Synthetic liquidity policy text"):
    stream = f"BT /F1 12 Tf 72 720 Td ({text}) Tj ET".encode()
    objs = [b"<< /Type /Catalog /Pages 2 0 R >>", b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
            b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents 4 0 R /Resources << /Font << /F1 5 0 R >> >> >>",
            b"<< /Length %d >>stream\n" % len(stream) + stream + b"\nendstream", b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>"]
    out, offs = b"%PDF-1.4\n", []
    for i, o in enumerate(objs, 1):
        offs.append(len(out))
        out += b"%d 0 obj\n" % i + o + b"\nendobj\n"
    x = len(out)
    out += b"xref\n0 %d\n0000000000 65535 f \n" % (len(objs) + 1) + b"".join(b"%010d 00000 n \n" % o for o in offs)
    out += b"trailer\n<< /Size %d /Root 1 0 R >>\nstartxref\n%d\n%%%%EOF" % (len(objs) + 1, x)
    return out


@pytest.mark.parametrize("name,data,needle", [("f.docx", _docx(), "0.35%"), ("f.xlsx", _xlsx(), "fee_bps: 35"), ("f.pdf", _pdf(), "liquidity"),
                                              ("f.csv", b"fund,fee\nA,10\n", "fee: 10")])
def test_parsers(name, data, needle):
    from retrieval.parse import parse_bytes
    from retrieval.validate import scan
    assert scan(name, data).ok
    _, sections = parse_bytes(name, data)
    assert needle in " ".join(t for _, t in sections)


@pytest.mark.critical
@pytest.mark.parametrize("name,data,reason", [
    ("x.md", b"X5O!P%@AP[4\\PZX54(P^)7CC)7}$EICAR-STANDARD-ANTIVIRUS-TEST-FILE!$H+H*", "malware"),
    ("x.exe", b"MZ\x90\x00", "extension"), ("x.pdf", b"MZ\x90", "executable"), ("x.pdf", b"%PDF-1.4 /JavaScript", "active content"),
    ("x.docx", b"not a zip", "zip"), ("x.md", b"\xff\xfe\x00bad", "UTF-8"), ("x.md", b"a" * (6 * 1024 * 1024), "5 MB")])
def test_malicious_files_rejected(name, data, reason):
    from retrieval.validate import scan
    r = scan(name, data)
    assert not r.ok and reason.lower() in r.reason.lower()


def test_macro_docx_rejected():
    import zipfile
    from retrieval.validate import scan
    b = io.BytesIO()
    with zipfile.ZipFile(b, "w") as z:
        z.writestr("word/vbaProject.bin", b"macro")
    assert "macro" in scan("m.docx", b.getvalue()).reason


def test_ingest_requires_metadata_and_upgrades_public_pii(st):
    from retrieval.ingest import IngestError, ingest_bytes
    with session_scope() as s:
        with pytest.raises(IngestError):
            ingest_bytes(s, "a.md", b"# t\nbody")
        meta = dict(id="DOC-T1", title="t", owner="o", version="1", classification="public", tenant="t-flo", access_groups=["public"],
                    effective_date="2026-01-01", jurisdiction="CA")
        d = ingest_bytes(s, "a.md", b"# Contact\nEmail jane.doe@example.com for details", meta)
        assert d.classification == "confidential"
        from backend.db import models as m
        s.query(m.DocChunk).filter_by(document_id="DOC-T1").delete()
        s.query(m.Document).filter_by(id="DOC-T1").delete()
    from retrieval.search import invalidate_index
    invalidate_index()


def test_scenarios_monotonic_and_bounds(st):
    from scenarios import engine as E
    a = Adapters(st, "t-flo")
    c = {r["scenario"]: r for r in E.compare(a)["scenarios"]}
    assert c["bull"]["aum_end"] > c["base"]["aum_end"] > c["bear"]["aum_end"] > c["severe_stress"]["aum_end"]
    zero = E.run(a, E.ScenarioInput(net_flows_pct=0, expected_equity_return=0, expected_bond_return=0, inflation=0, opex_growth=0), False)
    assert zero["results"]["aum_end"] < zero["results"]["aum_start"] + Decimal(1e9)  # only cash drift remains
    with pytest.raises(ValueError):
        E.run(a, E.ScenarioInput(equity_shock=-2.0))
    s = E.sensitivity(a)
    assert s["tornado"][0]["driver"] in ("equity_shock", "fee_compression_bps", "net_flows_pct")


def test_retirement_projection_math():
    from analytics.wealth import retirement_projection
    r = retirement_projection(Decimal(100000), Decimal(0), 1, 0.05, 0.0, retirement_years=1)
    assert r["balance_at_retirement"] == Decimal("105000.00")
    # with r=5%, i=0, N=1: W = B * r / (1 - 1/(1+r)) = 105000 * 1.05 = 110250 ; balance after = 0
    assert r["sustainable_first_year_withdrawal"] == Decimal("110250.00")
    assert r["path"][-1]["balance"] == Decimal("0.00")


def test_suitability_rules(st, engine):
    from analytics.wealth import suitability_check
    client = {"risk_profile": "Conservative", "kyc_last_review": date(2020, 1, 1)}
    r = suitability_check(client, [{"id": "PF", "volatility": 0.12, "top_weight": 0.2}], engine.bundle.suitability, date(2026, 9, 30))
    assert {f["rule"] for f in r["findings"]} == {"SUIT-VOL-BAND", "SUIT-SINGLE-ISSUER", "SUIT-KYC-STALE"}


@pytest.mark.parametrize("lang", ["en", "fr", "es", "hi", "ta", "ko"])
def test_communication_drafts_validate_in_every_language(ctx_for, a001_client, lang):
    from agents.comms import draft_communication, validate_translation
    d = draft_communication(ctx_for("u-advisor"), a001_client, "portfolio_review", lang)
    assert d["translation_validation"]["valid"], d["translation_validation"]
    assert d["compliance_precheck"]["passed"]
    assert not validate_translation("Value CAD 1,234.50 on 2026-09-30.", "Valeur CAD 1 234,50")["valid"]  # altered number is caught
