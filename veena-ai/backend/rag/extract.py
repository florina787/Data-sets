"""Text extraction for PDF, DOCX, TXT, Markdown, CSV and JSON."""

from __future__ import annotations

import csv
import io
import json
import re
from dataclasses import dataclass


@dataclass
class Page:
    number: int | None  # 1-based page for PDFs, None for flowing formats
    text: str


class ExtractionError(ValueError):
    pass


def _decode(data: bytes) -> str:
    for enc in ("utf-8-sig", "utf-16", "cp1252"):
        try:
            return data.decode(enc)
        except UnicodeDecodeError:
            continue
    return data.decode("utf-8", errors="replace")


def _pdf(data: bytes) -> list[Page]:
    import pymupdf

    try:
        doc = pymupdf.open(stream=data, filetype="pdf")
    except Exception as e:  # noqa: BLE001
        raise ExtractionError(f"not a readable PDF: {e}") from e
    if doc.needs_pass:
        raise ExtractionError("PDF is password protected")
    pages = [Page(i + 1, p.get_text("text")) for i, p in enumerate(doc)]
    doc.close()
    return pages


def _docx(data: bytes) -> list[Page]:
    import docx

    try:
        d = docx.Document(io.BytesIO(data))
    except Exception as e:  # noqa: BLE001
        raise ExtractionError(f"not a readable DOCX: {e}") from e
    parts = [p.text for p in d.paragraphs]
    for table in d.tables:
        for row in table.rows:
            parts.append(" | ".join(c.text.strip() for c in row.cells))
    return [Page(None, "\n".join(parts))]


def _markdown(data: bytes) -> list[Page]:
    from markdown_it import MarkdownIt

    text = _decode(data)
    # Keep the source text (headings and code stay meaningful) but drop HTML
    # blocks and inline HTML, which can hide content from the reader.
    tokens = MarkdownIt().parse(text)
    lines = text.splitlines()
    drop: set[int] = set()
    for t in tokens:
        if t.type == "html_block" and t.map:
            drop.update(range(t.map[0], t.map[1]))
    kept = [re.sub(r"<[^>]{1,200}>", "", line) for i, line in enumerate(lines) if i not in drop]
    return [Page(None, "\n".join(kept))]


def _csv(data: bytes) -> list[Page]:
    text = _decode(data)
    try:
        dialect = csv.Sniffer().sniff(text[:4096])
    except csv.Error:
        dialect = csv.excel
    rows = list(csv.reader(io.StringIO(text), dialect))
    if not rows:
        return [Page(None, "")]
    header, body = rows[0], rows[1:]
    # One line per record, "column: value", so each chunk is self-describing.
    lines = ["; ".join(f"{h}: {v}" for h, v in zip(header, r) if v.strip()) for r in body]
    return [Page(None, f"Columns: {', '.join(header)}\n" + "\n".join(lines))]


def _flatten(obj, prefix: str = "") -> list[str]:
    if isinstance(obj, dict):
        out = []
        for k, v in obj.items():
            out += _flatten(v, f"{prefix}.{k}" if prefix else str(k))
        return out
    if isinstance(obj, list):
        out = []
        for i, v in enumerate(obj):
            out += _flatten(v, f"{prefix}[{i}]")
        return out
    return [f"{prefix}: {obj}"]


def _json(data: bytes) -> list[Page]:
    try:
        obj = json.loads(_decode(data))
    except json.JSONDecodeError as e:
        raise ExtractionError(f"invalid JSON: {e}") from e
    return [Page(None, "\n".join(_flatten(obj)))]


EXTRACTORS = {
    ".pdf": _pdf,
    ".docx": _docx,
    ".txt": lambda d: [Page(None, _decode(d))],
    ".md": _markdown,
    ".markdown": _markdown,
    ".csv": _csv,
    ".json": _json,
}

MAGIC = {".pdf": b"%PDF", ".docx": b"PK\x03\x04"}


def extract(data: bytes, ext: str) -> list[Page]:
    ext = ext.lower()
    if ext not in EXTRACTORS:
        raise ExtractionError(f"unsupported type {ext}")
    if ext in MAGIC and not data.startswith(MAGIC[ext]):
        raise ExtractionError(f"file content does not match {ext}")
    if ext not in MAGIC and b"\x00" in data[:4096] and not data.startswith((b"\xff\xfe", b"\xfe\xff")):
        raise ExtractionError("binary content in a text file")
    return EXTRACTORS[ext](data)


def clean(text: str) -> str:
    text = text.replace("­", "").replace("﻿", "")
    text = re.sub(r"(\w)-\n(\w)", r"\1\2", text)  # de-hyphenate line breaks
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()
