import io
import json

import pytest

from backend.rag.chunk import chunk_pages
from backend.rag.extract import ExtractionError, Page, clean, extract
from backend.rag.service import lexical_score


def _pdf_bytes(pages: list[str]) -> bytes:
    import pymupdf

    doc = pymupdf.open()
    for text in pages:
        p = doc.new_page()
        p.insert_text((72, 72), text)
    data = doc.tobytes()
    doc.close()
    return data


def _docx_bytes(paras: list[str]) -> bytes:
    import docx

    d = docx.Document()
    for p in paras:
        d.add_paragraph(p)
    t = d.add_table(rows=1, cols=2)
    t.rows[0].cells[0].text = "Skill"
    t.rows[0].cells[1].text = "FAISS"
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def test_pdf_pages_numbered():
    pages = extract(_pdf_bytes(["First page about FAISS", "Second page about SQLite"]), ".pdf")
    assert [p.number for p in pages] == [1, 2]
    assert "SQLite" in pages[1].text


def test_docx_paragraphs_and_tables():
    text = extract(_docx_bytes(["Built a RAG pipeline"]), ".docx")[0].text
    assert "RAG pipeline" in text and "Skill | FAISS" in text


def test_csv_rows_self_describing():
    text = extract(b"name,role\nAsha,Engineer\nRavi,Architect\n", ".csv")[0].text
    assert "name: Asha; role: Engineer" in text


def test_json_flattened():
    text = extract(json.dumps({"project": {"name": "veena", "tags": ["ai", "local"]}}).encode(), ".json")[0].text
    assert "project.name: veena" in text and "project.tags[1]: local" in text


def test_markdown_drops_html():
    text = extract(b"# Title\n\n<div>hidden <b>html</b></div>\n\nVisible text", ".md")[0].text
    assert "Visible text" in text and "<div>" not in text


def test_magic_bytes_checked():
    with pytest.raises(ExtractionError):
        extract(b"not a pdf", ".pdf")
    with pytest.raises(ExtractionError):
        extract(b"\x00\x01binary", ".txt")


def test_clean_dehyphenates():
    assert clean("retrie-\nval   augmented\n\n\n\ngen") == "retrieval augmented\n\ngen"


def test_chunking_overlap_and_pages():
    long = "\n\n".join(f"Paragraph {i}. " + "word " * 60 for i in range(10))
    chunks = chunk_pages([Page(3, long)], size=500, overlap=100)
    assert len(chunks) > 3
    assert all(c.page == 3 for c in chunks)
    assert all(len(c.text) <= 700 for c in chunks)
    assert [c.ordinal for c in chunks] == list(range(len(chunks)))


def test_lexical_score():
    assert lexical_score("faiss vector index", "We used a FAISS vector index") > lexical_score("faiss", "nothing here")
